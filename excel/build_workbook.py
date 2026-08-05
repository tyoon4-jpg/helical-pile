"""
Builds HelicalPileDesign_CalcSheet.xlsx -- an Excel port of the
helical-pile-design Python calc engine (src/helical_pile_design/*.py),
pre-populated with the 12 m pole / medium-stiff-clay worked example from
tests/test_worked_example_12m.py.

Every formula in the workbook is native Excel -- no VBA/macros, and no
"enable iterative calculation" setting is required. See the Cover sheet
(built below) for the full design notes, especially how Step 5b's nonlinear
p-y solver is implemented (MMULT/MINVERSE matrix solves, with the Python
solver's fixed-point relaxation loop unrolled into sequential non-circular
"generation" blocks rather than a true circular reference -- see build_py_block()
and its docstring for why).

Requires: pip install openpyxl   (only needed to run this generator, not to
open the resulting .xlsx).

Run:      python build_workbook.py
Output:   ./HelicalPileDesign_CalcSheet.xlsx  (this directory)
"""
from __future__ import annotations

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.comments import Comment

# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")       # pale yellow -- editable
FILL_HEADER = PatternFill("solid", fgColor="0B72A6")       # step banner
FILL_SUBHEAD = PatternFill("solid", fgColor="D9E7F2")      # section banner
FILL_OUTPUT = PatternFill("solid", fgColor="E2EFDA")       # pale green -- key result
FILL_NOTE = PatternFill("solid", fgColor="FCE4D6")         # pale orange -- assumption/note

FONT_HEADER = Font(bold=True, size=13, color="FFFFFF")
FONT_SUBHEAD = Font(bold=True, size=10, color="1F3864")
FONT_LABEL = Font(size=10)
FONT_INPUT = Font(size=10, color="1F4E78", bold=True)
FONT_CALC = Font(size=10, color="000000")
FONT_OUTPUT = Font(size=10, bold=True)
FONT_CITE = Font(size=8, italic=True, color="7F7F7F")
FONT_NOTE = Font(size=8, italic=True, color="833C00")

THIN = Side(style="thin", color="BFBFBF")
BORDER_BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREEN_FONT = Font(color="006100", bold=True)
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006", bold=True)


def sheet_header(ws: Worksheet, row: int, title: str, span: int = 8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=title)
    c.font = FONT_HEADER
    c.fill = FILL_HEADER
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 22
    return row + 2


def subhead(ws: Worksheet, row: int, title: str, span: int = 8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=title)
    c.font = FONT_SUBHEAD
    c.fill = FILL_SUBHEAD
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return row + 1


def cite(ws: Worksheet, row: int, text: str, col: int = 1, span: int = 8):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    c = ws.cell(row=row, column=col, value=text)
    c.font = FONT_CITE
    c.alignment = Alignment(horizontal="left", wrap_text=True)
    return row + 1


def input_row(ws: Worksheet, row: int, label: str, value, unit: str = "", col_label=1, col_val=4, col_unit=5):
    lc = ws.cell(row=row, column=col_label, value=label)
    lc.font = FONT_LABEL
    vc = ws.cell(row=row, column=col_val, value=value)
    vc.font = FONT_INPUT
    vc.fill = FILL_INPUT
    vc.border = BORDER_BOX
    vc.number_format = "General"
    if unit:
        uc = ws.cell(row=row, column=col_unit, value=unit)
        uc.font = FONT_CITE
    return row + 1


def calc_row(ws: Worksheet, row: int, label: str, formula, unit: str = "", numfmt: str = "0.000",
             col_label=1, col_val=4, col_unit=5, bold_output=False):
    lc = ws.cell(row=row, column=col_label, value=label)
    lc.font = FONT_LABEL
    vc = ws.cell(row=row, column=col_val, value=formula)
    vc.font = FONT_OUTPUT if bold_output else FONT_CALC
    vc.fill = FILL_OUTPUT if bold_output else PatternFill(fill_type=None)
    vc.border = BORDER_BOX
    vc.number_format = numfmt
    if unit:
        uc = ws.cell(row=row, column=col_unit, value=unit)
        uc.font = FONT_CITE
    return row + 1


def passfail_row(ws: Worksheet, row: int, label: str, formula, col_label=1, col_val=4):
    lc = ws.cell(row=row, column=col_label, value=label)
    lc.font = Font(size=10, bold=True)
    vc = ws.cell(row=row, column=col_val, value=formula)
    vc.font = FONT_OUTPUT
    vc.border = BORDER_BOX
    addr = f"${get_column_letter(col_val)}${row}"
    ws.conditional_formatting.add(
        addr, FormulaRule(formula=[f'{addr}=TRUE'], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(
        addr, FormulaRule(formula=[f'{addr}=FALSE'], fill=RED_FILL, font=RED_FONT))
    return row + 1


def note_row(ws: Worksheet, row: int, text: str, span: int = 8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value="NOTE: " + text)
    c.font = FONT_NOTE
    c.fill = FILL_NOTE
    c.alignment = Alignment(horizontal="left", wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = max(15, 13 * (1 + len(text) // 90))
    return row + 1


def set_col_widths(ws: Worksheet, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


DEFAULT_WIDTHS = {"A": 40, "B": 3, "C": 3, "D": 14, "E": 10, "F": 3, "G": 40, "H": 14}

wb = Workbook()
wb.remove(wb.active)

# Cross-sheet cell references, filled in as each sheet is built.
R: dict[str, str] = {}

def A(sheet: str, row: int, col: int) -> str:
    """Cross-sheet absolute reference, e.g. A('1-Loads', 5, 4) -> "'1-Loads'!$D$5"."""
    return f"'{sheet}'!${get_column_letter(col)}${row}"


def a(row: int, col: int) -> str:
    """Same-sheet absolute reference."""
    return f"${get_column_letter(col)}${row}"


def plain(row: int, col: int) -> str:
    """Same-sheet reference with NO $ signs -- required for ArrayFormula's
    `ref` attribute (a structural range reference, not a formula operand;
    Excel silently refuses to open the file if it contains $ there)."""
    return f"{get_column_letter(col)}{row}"


# ===========================================================================
# COVER SHEET
# ===========================================================================
ws = wb.create_sheet("Cover")
set_col_widths(ws, {"A": 100})
r = 1
r = sheet_header(ws, r, "HELICAL PILE DESIGN -- EXCEL CALC SHEET", span=1)
ws.row_dimensions[1].height = 28
r += 1
lines = [
    ("Source", "Ported from the helical-pile-design Python calc engine "
     "(src/helical_pile_design/*.py). Every sheet maps 1:1 to the numbered "
     "step in the governing design procedure document, same as the source "
     "package's module structure."),
    ("Pre-loaded example", "All input cells (pale yellow) are pre-filled with the "
     "12 m pole / medium-stiff-clay worked example from "
     "tests/test_worked_example_12m.py -- change any yellow cell for a new project."),
    ("Units", "SI internally throughout (m, kN, kN*m, kPa), matching the source "
     "package convention. No unit conversions are hidden inside formulas."),
    ("Color key", "Pale YELLOW = editable input. Pale GREEN = key result. "
     "Pale ORANGE = assumption/limitation to confirm. GREEN/RED cell = pass/fail."),
    ("Step 5b -- IMPORTANT", "The p-y lateral solver (Step 5b) is a nonlinear "
     "finite-difference beam-on-Winkler-foundation solve. It is implemented with "
     "native array formulas (MMULT/MINVERSE) -- no macros, and NO 'enable iterative "
     "calculation' setting is needed. The Python solver's fixed-point relaxation "
     "loop is unrolled into 25-45 sequential 'generation' blocks stacked down the "
     "sheet, each depending only on the PREVIOUS generation's result (a one-way "
     "chain, not a circular reference), so Excel computes it in one ordinary pass. "
     "An earlier design nested the matrix solve inside a true self-referencing "
     "circular loop; tested in real Excel via COM automation, it proved unreliable "
     "-- a transient error during the first iterative pass would latch permanently "
     "and never resolve, a known rough edge of combining array formulas with "
     "circular references. The unrolled design avoids that failure mode entirely. "
     "The mesh here is coarsened (21 nodes / 20 elements for service+strength, "
     "11 nodes / 10 elements for the pushover sweep) vs. the Python reference's "
     "180 elements, to keep the workbook a manageable size. This was verified "
     "against the Python engine directly: at 20 elements, groundline deflection "
     "tracks the 180-element reference within ~4%, and pushover lambda_ult within "
     "the design procedure's expected 9.5-10 band -- comfortably inside this "
     "suite's own regression tolerances (15-30%). Treat this sheet as a calibrated "
     "engineering approximation, not a bit-exact reproduction; for final sealed "
     "calculations, cross-check against the Python package."),
    ("Two corrections vs. the source document", "(1) Section 11.5.1's Broms hand "
     "calculation (Hu~=27 kN) doesn't satisfy the document's own closure equation -- "
     "the correct root is Hu~=32.4 kN (FS~=11.8). (2) Section 3.1's shaft-sizing "
     "heuristic coefficient (40) is dimensionally inconsistent; this sheet uses the "
     "corrected coefficient (320), same as sizing.py. Neither changes the worked "
     "example's pass/fail conclusions -- both were governed by ample margins either way."),
    ("What this sheet does NOT do", "It does not duplicate calculation logic anywhere "
     "-- every formula traces to one function in the source package (cited on each "
     "sheet). It does not replace engineering judgment: cells marked "
     "'[ASSUMED -- validate]' (Kt, tau_ad, bolt thread-root modulus, etc.) still need "
     "project-specific verification, exactly as in the source app."),
]
for label, text in lines:
    r = subhead(ws, r, label, span=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
    c = ws.cell(row=r, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(size=10)
    ws.row_dimensions[r].height = 15 * (1 + len(text) // 95)
    r += 2

r = subhead(ws, r, "Step -> sheet map", span=1)
step_map = [
    "1  Loads                         -> 1-Loads",
    "2  Geotechnical parameterization -> 2-Geotech",
    "3  Preliminary sizing            -> 3-Sizing",
    "4  Axial capacity                -> 4-Axial",
    "5a Lateral -- Broms (prelim)     -> 5a-Broms",
    "5b Lateral -- p-y (final)        -> 5b-pySolver",
    "6  Structural (AISC 360)         -> 6-Structural",
    "7  Corrosion (75-yr life)        -> 7-Corrosion",
    "8  Pole-to-pile connection       -> 8-Connection",
    "9  Installation torque           -> 9-Torque",
    "10 QA/QC and load-test triggers  -> 10-QAQC",
    "   Pass/fail rollup              -> Summary",
]
for line in step_map:
    c = ws.cell(row=r, column=1, value=line)
    c.font = Font(size=10, name="Consolas")
    r += 1

print("cover sheet built")

# ===========================================================================
# STEP 1 -- LOADS
# ===========================================================================
SH1 = "1-Loads"
ws = wb.create_sheet(SH1)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 1 -- LOAD DETERMINATION")
r = cite(ws, r, "loads.py -- AASHTO LTS-6 wind pressure -> groundline V, M, P, Tz. "
                "Doc Section 1, eq. 1.1-1.4.")
r += 1

r = subhead(ws, r, "Pole geometry (A1-A3)")
r = input_row(ws, r, "Pole height, h", 12.0, "m"); h_row = r - 1
r = input_row(ws, r, "Base diameter, b_base", 0.25, "m"); bbase_row = r - 1
r = input_row(ws, r, "Top diameter, b_top", 0.10, "m"); btop_row = r - 1
r = input_row(ws, r, "Pole+arm projected area, A_pole", 2.10, "m^2"); apole_row = r - 1
r = input_row(ws, r, "Pole+arm weight", 1.50, "kN"); polewt_row = r - 1
r = input_row(ws, r, "Drag coefficient, Cd (round tapered)", 0.50, "-"); cd_row = r - 1
r += 1

r = subhead(ws, r, "Luminaire (A4-A5)")
r = input_row(ws, r, "Luminaire EPA (includes Cd)", 0.50, "m^2"); epa_row = r - 1
r = input_row(ws, r, "Luminaire weight", 0.25, "kN"); lumwt_row = r - 1
r = input_row(ws, r, "Mast-arm eccentricity, e_arm", 0.0, "m"); earm_row = r - 1
r += 1

r = subhead(ws, r, "Wind environment -- strength case (A6-A7)")
r = input_row(ws, r, "Basic wind speed, V (3-s gust, strength)", 51.4, "m/s"); vms_row = r - 1
r = input_row(ws, r, "Kz at luminaire height", 1.04, "-"); kztop_row = r - 1
r = input_row(ws, r, "Kz at pole area centroid", 1.00, "-"); kzc_row = r - 1
r = input_row(ws, r, "Directionality, Kd", 0.95, "-"); kd_row = r - 1
r = input_row(ws, r, "Gust factor, G", 1.14, "-"); g_row = r - 1
r = input_row(ws, r, "Service/strength pressure ratio (~10-yr MRI)", 0.44, "-"); spr_row = r - 1
r += 1

r = subhead(ws, r, "Load combination factors")
r = input_row(ws, r, "Reveal (rigid riser above grade)", 0.0, "m"); reveal_row = r - 1
r = input_row(ws, r, "Dead load factor", 1.0, "-"); dlf_row = r - 1
r = input_row(ws, r, "Wind load factor", 1.0, "-"); wlf_row = r - 1
r += 1

r = subhead(ws, r, "Calculated -- Strength case (Extreme I)")
r = cite(ws, r, "eq. 1.1: Pz = 0.613*Kz*Kd*G*V^2 [Pa]->kPa.  eq. 1.2: z_bar = (h/3)(b_base+2b_top)/(b_base+b_top).")
zbar_f = f"=({a(h_row,4)}/3)*({a(bbase_row,4)}+2*{a(btop_row,4)})/({a(bbase_row,4)}+{a(btop_row,4)})"
r = calc_row(ws, r, "Centroid height, z_bar", zbar_f, "m"); zbar_row = r - 1
pzc_f = f"=0.613*{a(kzc_row,4)}*{a(kd_row,4)}*{a(g_row,4)}*{a(vms_row,4)}^2/1000"
r = calc_row(ws, r, "Design pressure at z_bar, Pz(z_bar)", pzc_f, "kPa"); pzc_row = r - 1
pzt_f = f"=0.613*{a(kztop_row,4)}*{a(kd_row,4)}*{a(g_row,4)}*{a(vms_row,4)}^2/1000"
r = calc_row(ws, r, "Design pressure at h, Pz(h)", pzt_f, "kPa"); pzt_row = r - 1
fpole_f = f"={a(pzc_row,4)}*{a(cd_row,4)}*{a(apole_row,4)}"
r = calc_row(ws, r, "Pole+arm wind force, F_pole", fpole_f, "kN"); fpole_row = r - 1
flum_f = f"={a(pzt_row,4)}*{a(epa_row,4)}"
r = calc_row(ws, r, "Luminaire wind force, F_lum", flum_f, "kN"); flum_row = r - 1
V_f = f"={a(wlf_row,4)}*({a(fpole_row,4)}+{a(flum_row,4)})"
r = calc_row(ws, r, "Groundline shear, V", V_f, "kN", bold_output=True); V_row = r - 1
M_f = f"={a(wlf_row,4)}*({a(fpole_row,4)}*{a(zbar_row,4)}+{a(flum_row,4)}*{a(h_row,4)})+{a(V_row,4)}*{a(reveal_row,4)}"
r = calc_row(ws, r, "Groundline moment, M", M_f, "kN*m", bold_output=True); M_row = r - 1
P_f = f"={a(dlf_row,4)}*({a(polewt_row,4)}+{a(lumwt_row,4)})"
r = calc_row(ws, r, "Groundline axial, P", P_f, "kN", bold_output=True); P_row = r - 1
Tz_f = f"={a(wlf_row,4)}*{a(flum_row,4)}*{a(earm_row,4)}"
r = calc_row(ws, r, "Groundline torsion, Tz", Tz_f, "kN*m", bold_output=True); Tz_row = r - 1
e_f = f"=IF({a(V_row,4)}<>0,{a(M_row,4)}/{a(V_row,4)},\"inf\")"
r = calc_row(ws, r, "Load eccentricity, e = M/V", e_f, "m", bold_output=True); e_row = r - 1
r += 1

r = subhead(ws, r, "Calculated -- Service case (~10-yr MRI)")
r = cite(ws, r, "eq. 1.4 note: V_service = V_strength * SQRT(service_pressure_ratio) (Pz ~ V^2).")
vms_svc_f = f"={a(vms_row,4)}*SQRT({a(spr_row,4)})"
r = calc_row(ws, r, "Service wind speed, V_service", vms_svc_f, "m/s"); vmssvc_row = r - 1
pzc_svc_f = f"=0.613*{a(kzc_row,4)}*{a(kd_row,4)}*{a(g_row,4)}*{a(vmssvc_row,4)}^2/1000"
r = calc_row(ws, r, "Service pressure at z_bar", pzc_svc_f, "kPa"); pzcsvc_row = r - 1
pzt_svc_f = f"=0.613*{a(kztop_row,4)}*{a(kd_row,4)}*{a(g_row,4)}*{a(vmssvc_row,4)}^2/1000"
r = calc_row(ws, r, "Service pressure at h", pzt_svc_f, "kPa"); pztsvc_row = r - 1
fpole_svc_f = f"={a(pzcsvc_row,4)}*{a(cd_row,4)}*{a(apole_row,4)}"
r = calc_row(ws, r, "Service F_pole", fpole_svc_f, "kN"); fpolesvc_row = r - 1
flum_svc_f = f"={a(pztsvc_row,4)}*{a(epa_row,4)}"
r = calc_row(ws, r, "Service F_lum", flum_svc_f, "kN"); flumsvc_row = r - 1
Vsvc_f = f"={a(wlf_row,4)}*({a(fpolesvc_row,4)}+{a(flumsvc_row,4)})"
r = calc_row(ws, r, "Service shear, V_service", Vsvc_f, "kN", bold_output=True); Vsvc_row = r - 1
Msvc_f = f"={a(wlf_row,4)}*({a(fpolesvc_row,4)}*{a(zbar_row,4)}+{a(flumsvc_row,4)}*{a(h_row,4)})+{a(Vsvc_row,4)}*{a(reveal_row,4)}"
r = calc_row(ws, r, "Service moment, M_service", Msvc_f, "kN*m", bold_output=True); Msvc_row = r - 1
Psvc_f = f"={a(P_row,4)}"
r = calc_row(ws, r, "Service axial, P_service", Psvc_f, "kN"); Psvc_row = r - 1
esvc_f = f"=IF({a(Vsvc_row,4)}<>0,{a(Msvc_row,4)}/{a(Vsvc_row,4)},\"inf\")"
r = calc_row(ws, r, "Service eccentricity, e_service", esvc_f, "m"); esvc_row = r - 1

R.update({
    "L1.V": A(SH1, V_row, 4), "L1.M": A(SH1, M_row, 4), "L1.P": A(SH1, P_row, 4),
    "L1.Tz": A(SH1, Tz_row, 4), "L1.e": A(SH1, e_row, 4),
    "L1.Vsvc": A(SH1, Vsvc_row, 4), "L1.Msvc": A(SH1, Msvc_row, 4),
    "L1.esvc": A(SH1, esvc_row, 4),
    "L1.dlf": A(SH1, dlf_row, 4), "L1.wlf": A(SH1, wlf_row, 4),
})
ws.freeze_panes = "A4"
print("Step 1 (Loads) built")

# ===========================================================================
# STEP 2 -- GEOTECH
# ===========================================================================
SH2 = "2-Geotech"
ws = wb.create_sheet(SH2)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 2 -- GEOTECHNICAL PARAMETERIZATION")
r = cite(ws, r, "geotech.py -- SPT/CPT correlations -> design soil profile. Doc Section 2.")
r += 1

r = subhead(ws, r, "Reference correlations (informational -- prefer lab/CPT for design Su)")
r = input_row(ws, r, "SPT N60 (blows/ft), if available", 17, "-"); n60_row = r - 1
su_n60_f = f"=6*{a(n60_row,4)}"
r = calc_row(ws, r, "  Su from N60 = 6*N60 (eq. 2.1)", su_n60_f, "kPa")
r = input_row(ws, r, "CPT qt, if available", 0, "kPa"); qt_row = r - 1
r = input_row(ws, r, "In-situ sigma'v0 at CPT depth", 0, "kPa"); sv0_row = r - 1
r = input_row(ws, r, "CPT Nkt (12-16)", 14, "-"); nkt_row = r - 1
su_cpt_f = f"=IF({a(nkt_row,4)}=0,\"n/a\",({a(qt_row,4)}-{a(sv0_row,4)})/{a(nkt_row,4)})"
r = calc_row(ws, r, "  Su from CPT (eq. 2.1)", su_cpt_f, "kPa")
phi_f = f"=27.1+0.3*{a(n60_row,4)}-0.00054*{a(n60_row,4)}^2"
r = calc_row(ws, r, "  phi' from N60, Peck/Hanson (eq. 2.2, sand only)", phi_f, "deg")
r += 1

r = subhead(ws, r, "Design soil profile (governing clay layer within lateral influence zone)")
r = input_row(ws, r, "Design Su (governing layer)", 50.0, "kPa"); su_row = r - 1
lc = ws.cell(row=r, column=1, value="Basis")
lc.font = FONT_LABEL
bc = ws.cell(row=r, column=4, value="lab UU/CU")
bc.font = FONT_INPUT
bc.fill = FILL_INPUT
bc.border = BORDER_BOX
from openpyxl.worksheet.datavalidation import DataValidation
dv = DataValidation(type="list", formula1='"SPT correlation,lab UU/CU,CPT,lab + CPT"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(bc)
basis_row = r
r += 1
eps50_ref_f = (f'=IF({a(su_row,4)}<25,0.02,IF({a(su_row,4)}<=50,0.01,0.007))')
r = calc_row(ws, r, "  eps50 correlation vs. Su (reference, eq. 2.1)", eps50_ref_f, "-")
r = input_row(ws, r, "Design eps50 (override -- engineering judgment)", 0.007, "-"); eps50_row = r - 1
r = note_row(ws, r, "Doc worked example uses eps50=0.007 explicitly (stiff-clay band), overriding the "
                     "raw Su=50 correlation boundary value of 0.01 -- confirm against site data.")
r = input_row(ws, r, "Total unit weight, gamma (above GWT)", 18.5, "kN/m^3"); gamma_row = r - 1
r = input_row(ws, r, "Unit weight of water, gamma_w", 9.81, "kN/m^3"); gammaw_row = r - 1
gammasub_f = f"={a(gamma_row,4)}-{a(gammaw_row,4)}"
r = calc_row(ws, r, "Submerged unit weight, gamma_sub (below GWT)", gammasub_f, "kN/m^3"); gammasub_row = r - 1
r = input_row(ws, r, "Groundwater depth, GWT", 3.0, "m"); gwt_row = r - 1
r = input_row(ws, r, "Frost depth, d_f", 1.2, "m"); df_row = r - 1
r += 1
r = subhead(ws, r, "Step 10 QA/QC trigger flag")
lab_cpt_f = f'=IF({a(basis_row,4)}<>"SPT correlation",TRUE,FALSE)'
r = passfail_row(ws, r, "p-y params from lab/CPT (not SPT-only)?", lab_cpt_f)
lab_cpt_row = r - 1

R.update({
    "G2.Su": A(SH2, su_row, 4), "G2.eps50": A(SH2, eps50_row, 4),
    "G2.gamma": A(SH2, gamma_row, 4), "G2.gammasub": A(SH2, gammasub_row, 4),
    "G2.gwt": A(SH2, gwt_row, 4), "G2.df": A(SH2, df_row, 4),
    "G2.labcpt": A(SH2, lab_cpt_row, 4),
})
ws.freeze_panes = "A4"
print("Step 2 (Geotech) built")

# ===========================================================================
# STEP 3 -- SIZING
# ===========================================================================
SH3 = "3-Sizing"
ws = wb.create_sheet(SH3)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 3 -- PRELIMINARY PILE SIZING")
r = cite(ws, r, "sizing.py -- trial shaft/helix geometry + embedment screens. Doc Section 3.")
r += 1

r = subhead(ws, r, "Trial shaft diameter")
r = cite(ws, r, "Corrected heuristic (see sizing.py docstring): doc's own Section 3.1 coefficient (40) is "
                 "dimensionally inconsistent (gives 0.71 m, outside its own stated 168-324 mm typical range). "
                 "D_shaft >= SQRT(Mu/320) used here instead, calibrated to the stated range -- a rough "
                 "starting point only, refined by Steps 4-6.")
Mu_f = f"={R['L1.M']}"
r = calc_row(ws, r, "Design moment, Mu (from Step 1 strength case)", Mu_f, "kN*m"); mu_row = r - 1
dreq_f = f"=SQRT(MAX({a(mu_row,4)},0)/320)"
r = calc_row(ws, r, "Required D_shaft (heuristic)", dreq_f, "m"); dreq_row = r - 1
dnear_f = (f"=IF({a(dreq_row,4)}<=0.168,0.168,IF({a(dreq_row,4)}<=0.219,0.219,"
           f"IF({a(dreq_row,4)}<=0.273,0.273,0.324)))")
r = calc_row(ws, r, "Nearest standard OD (168/219/273/324 mm)", dnear_f, "m")
r = input_row(ws, r, "Selected shaft OD, D_shaft", 0.273, "m"); dshaft_row = r - 1
r = input_row(ws, r, "Nominal wall thickness, t_nominal", 9.3, "mm"); tnom_row = r - 1
r += 1

r = subhead(ws, r, "Trial helix configuration (doc Section 3.2)")
r = cite(ws, r, "Up to 3 helices supported. Leave diameter/depth blank for an unused plate.")
hdr = ws.cell(row=r, column=1, value="Helix #"); hdr.font = FONT_SUBHEAD
ws.cell(row=r, column=4, value="Diameter, D_h (m)").font = FONT_SUBHEAD
ws.cell(row=r, column=5, value="Depth, z (m)").font = FONT_SUBHEAD
r += 1
helix_rows = []
helix_defaults = [(0.450, 3.00), (0.450, 4.35), (None, None)]
for i, (dd, dz) in enumerate(helix_defaults, start=1):
    ws.cell(row=r, column=1, value=f"  Helix {i}").font = FONT_LABEL
    dc = ws.cell(row=r, column=4, value=dd)
    dc.font = FONT_INPUT; dc.fill = FILL_INPUT; dc.border = BORDER_BOX
    zc = ws.cell(row=r, column=5, value=dz)
    zc.font = FONT_INPUT; zc.fill = FILL_INPUT; zc.border = BORDER_BOX
    helix_rows.append(r)
    r += 1
diam_rng = f"{a(helix_rows[0],4)}:{a(helix_rows[-1],4)}"
depth_rng = f"{a(helix_rows[0],5)}:{a(helix_rows[-1],5)}"
r = input_row(ws, r, "Helix pitch", 0.075, "m"); pitch_row = r - 1
r = input_row(ws, r, "Shaft standoff below bottom helix", 0.15, "m"); standoff_row = r - 1
r += 1

r = subhead(ws, r, "Embedment (doc Section 3.3)")
topdepth_f = f"=MIN({depth_rng})"
r = calc_row(ws, r, "Top helix depth", topdepth_f, "m"); topdepth_row = r - 1
botdepth_f = f"=MAX({depth_rng})"
r = calc_row(ws, r, "Bottom helix depth", botdepth_f, "m"); botdepth_row = r - 1
tip_f = f"={a(botdepth_row,4)}+{a(standoff_row,4)}"
r = calc_row(ws, r, "Pile tip depth, L", tip_f, "m", bold_output=True); tip_row = r - 1
r = calc_row(ws, r, "Top helix diameter, D_h,top", None, "m"); dhtop_row = r - 1
ws.cell(row=dhtop_row, column=4).value = ArrayFormula(
    plain(dhtop_row, 4), f"=INDEX({diam_rng},MATCH({a(topdepth_row,4)},{depth_rng},0))")
frostreq_f = f"={R['G2.df']}+3*{a(dhtop_row,4)}"
r = calc_row(ws, r, "Frost-clearance requirement (top >= d_f+3*D_h)", frostreq_f, "m"); frostreq_row = r - 1
frostok_f = f"={a(topdepth_row,4)}>={a(frostreq_row,4)}"
r = passfail_row(ws, r, "Frost clearance OK?", frostok_f); frostok_row = r - 1
deepreq_f = f"=5*{a(dhtop_row,4)}"
r = calc_row(ws, r, "'Deep foundation' requirement (top >= 5*D_h)", deepreq_f, "m"); deepreq_row = r - 1
deepok_f = f"={a(topdepth_row,4)}>={a(deepreq_row,4)}"
r = passfail_row(ws, r, "Deep-foundation mode OK?", deepok_f); deepok_row = r - 1
lratio_f = f"={a(tip_row,4)}/{a(dshaft_row,4)}"
r = calc_row(ws, r, "L / D_shaft (informational -- verify vs. Step 5 required length)", lratio_f, "-")
allpass_f = f"=AND({a(frostok_row,4)},{a(deepok_row,4)})"
r = passfail_row(ws, r, "Step 3 embedment screens -- ALL PASS?", allpass_f); allpass3_row = r - 1

R.update({
    "S3.Dshaft": A(SH3, dshaft_row, 4), "S3.tnom": A(SH3, tnom_row, 4),
    "S3.diam_rng": f"'{SH3}'!{diam_rng}", "S3.depth_rng": f"'{SH3}'!{depth_rng}",
    "S3.topdepth": A(SH3, topdepth_row, 4), "S3.botdepth": A(SH3, botdepth_row, 4),
    "S3.tip": A(SH3, tip_row, 4), "S3.allpass": A(SH3, allpass3_row, 4),
    "S3.helix1diam": A(SH3, helix_rows[0], 4), "S3.helix2diam": A(SH3, helix_rows[1], 4),
})
ws.freeze_panes = "A4"
print("Step 3 (Sizing) built")

# ===========================================================================
# STEP 4 -- AXIAL
# ===========================================================================
SH4 = "4-Axial"
ws = wb.create_sheet(SH4)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 4 -- AXIAL CAPACITY (COMPRESSION & TENSION)")
r = cite(ws, r, "axial.py -- individual bearing + cylindrical shear (lesser governs); "
                 "frost-uplift demand. Doc Section 4. FS_required = 2.0.")
r += 1

r = subhead(ws, r, "Shared inputs")
r = calc_row(ws, r, "Shaft OD, D_shaft (Step 3)", f"={R['S3.Dshaft']}", "m"); dshaft_row = r - 1
r = calc_row(ws, r, "Design Su (Step 2)", f"={R['G2.Su']}", "kPa"); su_row = r - 1
r = input_row(ws, r, "Effective shaft-adhesion length, L_eff (frost zone to top helix, doc-specific)", 1.35, "m")
leff_row = r - 1
r = note_row(ws, r, "L_eff is a caller-supplied parameter in axial.py (not auto-derived from helix depth) -- "
                     "the worked example's 1.35 m reflects the doc's specific adhesion-zone definition; "
                     "confirm against project geometry rather than assuming L_eff = top-helix-depth minus d_f.")
alpha_formula_f = (f"=IF({a(su_row,4)}<=25,1,IF({a(su_row,4)}>=70,0.5,"
                    f"1+({a(su_row,4)}-25)*(0.5-1)/(70-25)))")
r = calc_row(ws, r, "  alpha (API linear-ramp correlation, reference)", alpha_formula_f, "-")
r = input_row(ws, r, "Alpha override (blank cell -> use correlation above)", 0.5, "-"); alphaov_row = r - 1
alpha_f = f"=IF({a(alphaov_row,4)}=\"\",{alpha_formula_f.lstrip('=')},{a(alphaov_row,4)})"
r = calc_row(ws, r, "Alpha used", alpha_f, "-"); alpha_row = r - 1
r += 1

r = subhead(ws, r, "4.1 Individual bearing method (clay)")
r = cite(ws, r, "Qu = sum(A_h,i * 9*Su) + alpha*Su*pi*D_shaft*L_eff.  A_h = pi/4*(D_h^2-D_shaft^2).  eq. 4.1.")
helix_area_terms = []
for i, hrow in enumerate(helix_rows, start=1):
    area_f = (f"=IF('{SH3}'!{a(hrow,4)}=\"\",0,PI()/4*('{SH3}'!{a(hrow,4)}^2-{a(dshaft_row,4)}^2))")
    r = calc_row(ws, r, f"  Helix {i} net area, A_h", area_f, "m^2", numfmt="0.0000")
    helix_area_terms.append(a(r - 1, 4))
    r = calc_row(ws, r, f"  Helix {i} bearing = A_h*9*Su", f"={helix_area_terms[-1]}*9*{a(su_row,4)}", "kN")
    helix_area_terms[-1] = a(r - 1, 4)  # now points at the bearing-force row, not area
helix_bearing_sum = "+".join(helix_area_terms)
shaft_term_f = f"={a(alpha_row,4)}*{a(su_row,4)}*PI()*{a(dshaft_row,4)}*{a(leff_row,4)}"
r = calc_row(ws, r, "Shaft adhesion term = alpha*Su*pi*D_shaft*L_eff", shaft_term_f, "kN"); shaftterm_row = r - 1
qu_ind_f = f"=({helix_bearing_sum})+{a(shaftterm_row,4)}"
r = calc_row(ws, r, "Qu, individual bearing (compression)", qu_ind_f, "kN", bold_output=True); quind_row = r - 1
r += 1

r = subhead(ws, r, "4.2 Cylindrical shear method")
r = cite(ws, r, "Qu = A_h,bottom*9*Su + Su*pi*D_h,avg*L_c + alpha*Su*pi*D_shaft*L_eff. "
                 "Governs when helices are closely spaced. eq. 4.2.")
dhbottom_f = f"=INDEX({R['S3.diam_rng']},MATCH({R['S3.botdepth']},{R['S3.depth_rng']},0))"
r = calc_row(ws, r, "Bottom helix diameter, D_h,bottom", None, "m")
ws.cell(row=r - 1, column=4).value = ArrayFormula(plain(r - 1, 4), dhbottom_f)
dhbottom_row = r - 1
dhavg_f = f"=AVERAGE({R['S3.diam_rng']})"
r = calc_row(ws, r, "Average helix diameter, D_h,avg", dhavg_f, "m"); dhavg_row = r - 1
lc_f = f"={R['S3.botdepth']}-{R['S3.topdepth']}"
r = calc_row(ws, r, "Cylinder length between helices, L_c", lc_f, "m"); lc_row = r - 1
bottombearing_f = f"=PI()/4*({a(dhbottom_row,4)}^2-{a(dshaft_row,4)}^2)*9*{a(su_row,4)}"
r = calc_row(ws, r, "Bottom-helix bearing", bottombearing_f, "kN"); bottombearing_row = r - 1
cylshear_f = f"={a(su_row,4)}*PI()*{a(dhavg_row,4)}*{a(lc_row,4)}"
r = calc_row(ws, r, "Cylinder shear", cylshear_f, "kN"); cylshear_row = r - 1
qu_cyl_f = f"={a(bottombearing_row,4)}+{a(cylshear_row,4)}+{a(shaftterm_row,4)}"
r = calc_row(ws, r, "Qu, cylindrical shear (compression)", qu_cyl_f, "kN", bold_output=True); qucyl_row = r - 1
r += 1

r = subhead(ws, r, "4.3 Frost-heave uplift demand")
r = cite(ws, r, "T_frost = tau_ad * pi * D_shaft * d_f.  eq. 4.3.")
r = input_row(ws, r, "Adfreeze bond stress, tau_ad [ASSUMED -- validate]", 20.0, "kPa"); tauad_row = r - 1
tfrost_f = f"={a(tauad_row,4)}*PI()*{a(dshaft_row,4)}*{R['G2.df']}"
r = calc_row(ws, r, "Frost-uplift demand, T_frost", tfrost_f, "kN", bold_output=True); tfrost_row = r - 1
r += 1

r = subhead(ws, r, "4.1 Uplift (tension) resistance")
r = cite(ws, r, "Tu = sum(A_h,i*9*Su) [+ optional shaft adhesion below d_f].")
tu_ind_f = f"=({helix_bearing_sum})"
r = calc_row(ws, r, "Tu, individual bearing (tension, helix only)", tu_ind_f, "kN", bold_output=True); tuind_row = r - 1
r = input_row(ws, r, "Tu, cylindrical-shear tension [not defined by source module -- "
                      "enter separately if a project-specific method applies, else leave = Tu,individual]",
              None, "kN")
tucyl_input_row = r - 1
tucyl_f = f"=IF({a(tucyl_input_row,4)}=\"\",{a(tuind_row,4)},{a(tucyl_input_row,4)})"
r = calc_row(ws, r, "Tu, cylindrical (used)", tucyl_f, "kN"); tucyl_row = r - 1
r += 1

r = subhead(ws, r, "4.4 Acceptance -- FS_required = 2.0")
r = input_row(ws, r, "Axial demand override, P_max (compression) [doc worked value includes "
                      "adapter hardware + 1.25 DL factor not modeled in Step 1]", 2.6, "kN")
pmax_row = r - 1
r = calc_row(ws, r, "  (for reference: Step 1 P_strength)", f"={R['L1.P']}", "kN")
qgov_c_f = f"=MIN({a(quind_row,4)},{a(qucyl_row,4)})"
r = calc_row(ws, r, "Governing Qu (compression)", qgov_c_f, "kN", bold_output=True); qgovc_row = r - 1
methodc_f = f'=IF({a(quind_row,4)}<={a(qucyl_row,4)},"individual bearing","cylindrical shear")'
r = calc_row(ws, r, "  Governing method", methodc_f, "-", numfmt="General")
fsc_f = f"=IF({a(pmax_row,4)}>0,{a(qgovc_row,4)}/{a(pmax_row,4)},\"inf\")"
r = calc_row(ws, r, "FS, compression", fsc_f, "-", numfmt="0.00", bold_output=True); fsc_row = r - 1
passc_f = f"={a(fsc_row,4)}>=2"
r = passfail_row(ws, r, "Compression passes (FS>=2.0)?", passc_f); passc_row = r - 1
r += 1
qgov_t_f = f"=MIN({a(tuind_row,4)},{a(tucyl_row,4)})"
r = calc_row(ws, r, "Governing Tu (tension)", qgov_t_f, "kN", bold_output=True); qgovt_row = r - 1
fst_f = f"=IF({a(tfrost_row,4)}>0,{a(qgovt_row,4)}/{a(tfrost_row,4)},\"inf\")"
r = calc_row(ws, r, "FS, tension (vs. frost uplift demand)", fst_f, "-", numfmt="0.00", bold_output=True); fst_row = r - 1
passt_f = f"={a(fst_row,4)}>=2"
r = passfail_row(ws, r, "Tension passes (FS>=2.0)?", passt_f); passt_row = r - 1

R.update({
    "A4.passc": A(SH4, passc_row, 4), "A4.passt": A(SH4, passt_row, 4),
    "A4.fsc": A(SH4, fsc_row, 4), "A4.fst": A(SH4, fst_row, 4),
    "A4.pmax": A(SH4, pmax_row, 4), "A4.tfrost": A(SH4, tfrost_row, 4),
    "A4.yratio_note": None,
})
ws.freeze_panes = "A4"
print("Step 4 (Axial) built")

# ===========================================================================
# STEP 5a -- BROMS (preliminary lateral, closed-form via bisection ladder)
# ===========================================================================
SH5A = "5a-Broms"
ws = wb.create_sheet(SH5A)
set_col_widths(ws, {"A": 34, "B": 3, "C": 14, "D": 14, "E": 14, "F": 3,
                     "G": 14, "H": 14, "I": 14, "J": 14})
r = 1
r = sheet_header(ws, r, "STEP 5a -- BROMS METHOD (PRELIMINARY LATERAL CHECK)", span=10)
r = cite(ws, r, "lateral_py.broms_ultimate_clay_kN -- free-head short rigid pile in clay. "
                 "Doc eq. 5.2. FS_required = 2.0.", span=10)
r += 1
r = subhead(ws, r, "Inputs (linked)", span=10)
r = calc_row(ws, r, "Su", f"={R['G2.Su']}", "kPa"); su5_row = r - 1
r = calc_row(ws, r, "D_shaft", f"={R['S3.Dshaft']}", "m"); dshaft5_row = r - 1
r = calc_row(ws, r, "Embedded length, L (pile tip depth)", f"={R['S3.tip']}", "m"); l5_row = r - 1
r = calc_row(ws, r, "Load eccentricity, e = M/V (Step 1 strength)", f"={R['L1.e']}", "m"); e5_row = r - 1
r = calc_row(ws, r, "Frost depth, d_f", f"={R['G2.df']}", "m"); df5_row = r - 1
r = calc_row(ws, r, "Design shear demand, V (Step 1 strength)", f"={R['L1.V']}", "kN"); v5_row = r - 1
r += 1
r = subhead(ws, r, "Closure-equation setup", span=10)
r = cite(ws, r, "z0 = MAX(1.5*D_shaft, d_f).  Available length must satisfy L = z0 + f + g.", span=10)
z0_f = f"=MAX(1.5*{a(dshaft5_row,4)},{a(df5_row,4)})"
r = calc_row(ws, r, "Exclusion depth, z0", z0_f, "m"); z0_row = r - 1
avail_f = f"={a(l5_row,4)}-{a(z0_row,4)}"
r = calc_row(ws, r, "Available length, avail = L - z0", avail_f, "m"); avail_row = r - 1
r = note_row(ws, r, "If 'avail' <= 0, the embedded length does not clear the exclusion depth -- "
                     "lengthen the pile before proceeding (matches the ValueError raised by "
                     "broms_ultimate_clay_kN in the source module).", span=10)
r += 1

r = subhead(ws, r, "Bisection solve for Hu (closure equation g(Hu) = f + g - avail = 0)", span=10)
r = cite(ws, r, "f = Hu/(9*Su*D_shaft).  Mmax = Hu*(e+z0+0.5f).  g = SQRT(MAX(Mmax,0)/(2.25*Su*D_shaft)). "
                 "60-step bisection ladder -- plain formulas, no circular references, no macros.", span=10)
hdr_row = r
headers = ["iter", "lo (kN)", "hi (kN)", "mid = Hu trial", "f", "Mmax", "g", "g(Hu)=f+g-avail"]
for ci, htext in enumerate(headers, start=2):
    c = ws.cell(row=hdr_row, column=ci, value=htext)
    c.font = FONT_SUBHEAD
    c.fill = FILL_SUBHEAD
r += 1
N_BISECT = 60
first_data_row = r
for k in range(N_BISECT):
    row = r + k
    ws.cell(row=row, column=2, value=k + 1)
    if k == 0:
        ws.cell(row=row, column=3, value=0.000001)   # lo seed
        ws.cell(row=row, column=4, value=100000)      # hi seed -- generous upper bound
    else:
        prev = row - 1
        ws.cell(row=row, column=3, value=f"=IF({a(prev,9)}>0,{a(prev,3)},{a(prev,5)})")
        ws.cell(row=row, column=4, value=f"=IF({a(prev,9)}>0,{a(prev,5)},{a(prev,4)})")
    ws.cell(row=row, column=5, value=f"=({a(row,3)}+{a(row,4)})/2")            # mid
    ws.cell(row=row, column=6, value=f"={a(row,5)}/(9*{a(su5_row,4)}*{a(dshaft5_row,4)})")  # f
    ws.cell(row=row, column=7,
            value=f"={a(row,5)}*({a(e5_row,4)}+{a(z0_row,4)}+0.5*{a(row,6)})")  # Mmax
    ws.cell(row=row, column=8,
            value=f"=SQRT(MAX({a(row,7)},0)/(2.25*{a(su5_row,4)}*{a(dshaft5_row,4)}))")  # g
    ws.cell(row=row, column=9, value=f"={a(row,6)}+{a(row,8)}-{a(avail_row,4)}")  # g(Hu)
    for ci in range(2, 10):
        ws.cell(row=row, column=ci).number_format = "0.000000"
last_row = r + N_BISECT - 1
r = last_row + 1
r += 1

r = subhead(ws, r, "Result", span=10)
hu_f = f"={a(last_row,5)}"
r = calc_row(ws, r, "Hu, Broms ultimate lateral resistance", hu_f, "kN", numfmt="0.00", bold_output=True)
hu_row = r - 1
fs_f = f"={a(hu_row,4)}/{a(v5_row,4)}"
r = calc_row(ws, r, "FS, Broms", fs_f, "-", numfmt="0.00", bold_output=True); fsbroms_row = r - 1
pass_f = f"={a(fsbroms_row,4)}>=2"
r = passfail_row(ws, r, "Broms preliminary check passes (FS>=2.0)?", pass_f); passbroms_row = r - 1

R.update({
    "B5.Hu": A(SH5A, hu_row, 4), "B5.FS": A(SH5A, fsbroms_row, 4), "B5.pass": A(SH5A, passbroms_row, 4),
})
ws.freeze_panes = "A4"
print("Step 5a (Broms) built")


# ===========================================================================
# STEP 5b -- p-y SOLVER helpers (coarse FD beam-on-Winkler, native formulas)
# ===========================================================================
# Element stiffness matrix (Euler-Bernoulli beam, length h, constant EI):
#   k_elem = [[a, b,-a, b],
#             [b, c,-b, d],
#             [-a,-b, a,-b],
#             [b, d,-b, c]]
# where a=12EI/h^3, b=6EI/h^2, c=4EI/h, d=2EI/h. Symbolic assembly below
# mirrors lateral_py.solve_py's `K_beam[np.ix_(idx, idx)] += k_elem` loop
# exactly, but keeps every entry as a linear combination of the four named
# coefficient cells so Excel can recompute it live if D/t/E/L change.
_KELEM = [
    [("a", 1), ("b", 1), ("a", -1), ("b", 1)],
    [("b", 1), ("c", 1), ("b", -1), ("d", 1)],
    [("a", -1), ("b", -1), ("a", 1), ("b", -1)],
    [("b", 1), ("d", 1), ("b", -1), ("c", 1)],
]


def symbolic_beam_K(n_elements: int) -> dict:
    """Returns {(gi, gj): {coef: total_multiplier}} for the assembled global
    beam stiffness, 0-indexed dof (ndof = 2*(n_elements+1))."""
    K: dict = {}
    for e in range(n_elements):
        idx = [2 * e, 2 * e + 1, 2 * e + 2, 2 * e + 3]
        for i in range(4):
            for j in range(4):
                coef, mult = _KELEM[i][j]
                gi, gj = idx[i], idx[j]
                cell = K.setdefault((gi, gj), {})
                cell[coef] = cell.get(coef, 0) + mult
    return K


def render_K_formula(coefs: dict, coef_addr: dict, ks_addr: str | None) -> str:
    terms = []
    for coef, mult in coefs.items():
        if mult == 0:
            continue
        ref = coef_addr[coef]
        if mult == 1:
            terms.append(ref)
        elif mult == -1:
            terms.append(f"-{ref}")
        else:
            terms.append(f"{mult}*{ref}")
    if ks_addr:
        terms.append(ks_addr)
    if not terms:
        # Some (gi,gj) beam-assembly coefficients cancel to exactly zero (e.g.
        # adjacent-element contributions to the same dof pair). Returning the
        # bare string "0" here would make openpyxl write a TEXT cell (since it
        # doesn't start with "="), and a single text cell anywhere inside a
        # MINVERSE input range poisons the whole array with #VALUE! -- this
        # was the root cause of an earlier hard-to-find bug. "=0" is an actual
        # formula (numeric), matching every other branch's leading "=".
        return "=0"
    return "=" + "+".join(terms)


def build_py_block(
    ws: Worksheet, row0: int, col0: int, n_elements: int, n_generations: int, label: str,
    V_ref: str, M_ref: str, D_ref: str, t_mm_ref: str, L_ref: str,
    su_ref: str, gamma_ref: str, gammasub_ref: str, gwt_ref: str,
    eps50_ref: str, frost_ref: str,
) -> dict:
    """Lays out one coarse nonlinear p-y solve block (service/strength/pushover
    load case) as native Excel formulas.

    Design note: an earlier version nested the MMULT/MINVERSE matrix solve
    inside a self-referencing (circular) cell per node, relying on Excel's
    iterative-calculation engine to converge it -- mirroring lateral_py.solve_py's
    Python loop directly. That was tested in real Excel (via COM automation)
    and found to be unreliable: once any transient state during the first
    iterative pass produces an error (#VALUE!) inside the circular web, the
    error latches permanently (error + number = error, forever), because
    Excel's iterative engine has no way to "recover" from an error the way it
    recovers numerically. This is a known rough edge of combining array
    formulas with circular references.

    This version instead UNROLLS the fixed-point relaxation loop into
    `n_generations` sequential blocks stacked down the sheet -- generation g's
    K matrix depends only on generation (g-1)'s y-vector, a strictly one-way
    dependency (a DAG, not a cycle). No circular references, no iterative
    calculation setting required, fully deterministic. Verified in Python
    against the reference solver that ~20-30 generations at this mesh size
    converges to within engineering tolerance (residual < 1e-5 m) well before
    the generation budget used here.
    """
    nn = n_elements + 1
    ndof = 2 * nn
    r = row0
    r = subhead(ws, r, label, span=13)
    r += 1

    # --- parameters panel -------------------------------------------------
    def prm(name, formula, unit=""):
        nonlocal r
        ws.cell(row=r, column=col0, value=name).font = FONT_LABEL
        vc = ws.cell(row=r, column=col0 + 1, value=formula)
        vc.font = FONT_CALC
        vc.number_format = "0.00000"
        if unit:
            ws.cell(row=r, column=col0 + 2, value=unit).font = FONT_CITE
        addr = a(r, col0 + 1)
        r += 1
        return addr

    V_addr = prm("V (kN)", f"={V_ref}", "kN")
    M_addr = prm("M (kN*m)", f"={M_ref}", "kN*m")
    D_addr = prm("D_shaft", f"={D_ref}", "m")
    t_addr = prm("t (corroded)", f"={t_mm_ref}/1000", "m")
    E_addr = prm("E, steel", "=200000000", "kPa")
    I_addr = prm("I (pipe)", f"=PI()/64*({D_addr}^4-({D_addr}-2*{t_addr})^4)", "m^4")
    EI_addr = prm("EI", f"={E_addr}*{I_addr}", "kN*m^2")
    L_addr = prm("L, embedded length", f"={L_ref}", "m")
    n_addr = prm("n_elements", n_elements, "-")
    h_addr = prm("h = L/n", f"={L_addr}/{n_addr}", "m")
    a_addr = prm("a = 12EI/h^3", f"=12*{EI_addr}/{h_addr}^3")
    b_addr = prm("b = 6EI/h^2", f"=6*{EI_addr}/{h_addr}^2")
    c_addr = prm("c = 4EI/h", f"=4*{EI_addr}/{h_addr}")
    d_addr = prm("d = 2EI/h", f"=2*{EI_addr}/{h_addr}")
    su_addr = prm("Su", f"={su_ref}", "kPa")
    gamma_addr = prm("gamma", f"={gamma_ref}", "kN/m3")
    gammasub_addr = prm("gamma_sub", f"={gammasub_ref}", "kN/m3")
    gwt_addr = prm("GWT depth", f"={gwt_ref}", "m")
    eps50_addr = prm("eps50", f"={eps50_ref}", "-")
    frost_addr = prm("frost depth", f"={frost_ref}", "m")
    J_addr = prm("J (Matlock)", 0.5, "-")
    y50_addr = prm("y50 = 2.5*eps50*D", f"=2.5*{eps50_addr}*{D_addr}", "m")
    n_gen_addr = prm("n_generations (unrolled iterations)", n_generations, "-")
    coef_addr = {"a": a_addr, "b": b_addr, "c": c_addr, "d": d_addr}

    r += 1
    # --- static node geometry (independent of y, computed once) -----------
    node_col0 = col0 + 3
    headers = ["i", "z (m)", "sigma_v (kPa)", "Np", "pu (kN/m)", "trib (m)"]
    for j, htxt in enumerate(headers):
        c = ws.cell(row=r, column=node_col0 + j, value=htxt)
        c.font = FONT_SUBHEAD
        c.fill = FILL_SUBHEAD
    r += 1
    static_row0 = r
    SC = {name: node_col0 + j for j, name in enumerate(["i", "z", "sv", "Np", "pu", "trib"])}
    for i in range(nn):
        row = static_row0 + i
        ws.cell(row=row, column=SC["i"], value=i)
        i_c = a(row, SC["i"])
        ws.cell(row=row, column=SC["z"], value=f"={i_c}*{h_addr}")
        z_c = a(row, SC["z"])
        ws.cell(row=row, column=SC["sv"],
                value=f"=IF({z_c}<={gwt_addr},{gamma_addr}*{z_c},"
                      f"{gamma_addr}*{gwt_addr}+{gammasub_addr}*({z_c}-{gwt_addr}))")
        sv_c = a(row, SC["sv"])
        ws.cell(row=row, column=SC["Np"], value=f"=MIN(3+{sv_c}/{su_addr}+{J_addr}*{z_c}/{D_addr},9)")
        np_c = a(row, SC["Np"])
        ws.cell(row=row, column=SC["pu"], value=f"={np_c}*{su_addr}*{D_addr}")
        ws.cell(row=row, column=SC["trib"],
                value=f"=IF(OR({i_c}=0,{i_c}={n_addr}),{h_addr}/2,{h_addr})")
        for col in ("z", "sv", "Np", "pu", "trib"):
            ws.cell(row=row, column=SC[col]).number_format = "0.000000"
    static_row_last = static_row0 + nn - 1
    r = static_row_last + 2

    # --- shared force vector (unchanged across generations) ----------------
    Fcol = node_col0
    Frow0 = r
    for gi in range(ndof):
        ws.cell(row=Frow0 + gi, column=Fcol,
                value=(f"={V_addr}" if gi == 0 else (f"={M_addr}" if gi == 1 else 0)))
        ws.cell(row=Frow0 + gi, column=Fcol).number_format = "0.000"
    F_range = f"{a(Frow0, Fcol)}:{a(Frow0 + ndof - 1, Fcol)}"
    r = Frow0 + ndof + 2

    # --- generation 0 seed --------------------------------------------------
    seed_col = node_col0 + 1
    seed_row0 = r
    ws.cell(row=r - 1, column=col0, value=f"y seed (gen 0)").font = FONT_CITE
    for i in range(nn):
        ws.cell(row=seed_row0 + i, column=seed_col, value=0.00001).number_format = "0.000000"
    r = seed_row0 + nn + 1
    prev_y_row0, prev_y_col = seed_row0, seed_col

    ks_col, y_col, theta_col = node_col0, node_col0 + 1, node_col0 + 2
    Kcol0 = node_col0
    Ksym = symbolic_beam_K(n_elements)

    final_y_row0 = final_theta_row0 = None
    for g in range(1, n_generations + 1):
        gen_row0 = r
        ws.cell(row=gen_row0 - 1, column=col0, value=f"gen {g}").font = FONT_CITE
        for i in range(nn):
            row = gen_row0 + i
            prev_y_cell = a(prev_y_row0 + i, prev_y_col)
            z_c = a(static_row0 + i, SC["z"])
            pu_c = a(static_row0 + i, SC["pu"])
            trib_c = a(static_row0 + i, SC["trib"])
            yeff_e = f"IF(ABS({prev_y_cell})>0.000000001,{prev_y_cell},0.000000001)"
            ay_e = f"ABS({yeff_e})"
            pmag_e = (f"IF({z_c}<={frost_addr},0,IF({ay_e}>=8*{y50_addr},{pu_c},"
                      f"0.5*{pu_c}*({ay_e}/{y50_addr})^(1/3)))")
            ks_f = f"=IF({ay_e}=0,0,{pmag_e}/{ay_e}*{trib_c})"
            ws.cell(row=row, column=ks_col, value=ks_f).number_format = "0.000000"
        Krow0 = gen_row0 + nn + 1
        ks_cell_by_dof = {2 * i: a(gen_row0 + i, ks_col) for i in range(nn)}
        for gi in range(ndof):
            for gj in range(ndof):
                coefs = Ksym.get((gi, gj))
                ks_addr_ = ks_cell_by_dof.get(gi) if gi == gj else None
                cell = ws.cell(row=Krow0 + gi, column=Kcol0 + gj)
                if coefs:
                    cell.value = render_K_formula(coefs, coef_addr, ks_addr_)
                elif ks_addr_:
                    cell.value = f"={ks_addr_}"
                else:
                    cell.value = 0
                cell.number_format = "0.000"
        Kcol_last = Kcol0 + ndof - 1
        Ucol = Kcol_last + 1
        u_formula = f"=MMULT(MINVERSE({a(Krow0,Kcol0)}:{a(Krow0+ndof-1,Kcol_last)}),{F_range})"
        ws.cell(row=Krow0, column=Ucol).value = ArrayFormula(
            ref=f"{plain(Krow0, Ucol)}:{plain(Krow0 + ndof - 1, Ucol)}", text=u_formula)
        for gi in range(ndof):
            ws.cell(row=Krow0 + gi, column=Ucol).number_format = "0.000000"

        for i in range(nn):
            row = gen_row0 + i
            u_y = a(Krow0 + 2 * i, Ucol)
            u_th = a(Krow0 + 2 * i + 1, Ucol)
            prev_y_cell = a(prev_y_row0 + i, prev_y_col)
            ws.cell(row=row, column=y_col, value=f"=0.5*{prev_y_cell}+0.5*{u_y}").number_format = "0.000000"
            ws.cell(row=row, column=theta_col, value=f"={u_th}").number_format = "0.000000"

        r = Krow0 + ndof + 2
        prev_y_row0, prev_y_col = gen_row0, y_col
        final_y_row0, final_theta_row0 = gen_row0, gen_row0

    # --- element table (for M_node), using the FINAL generation's y/theta --
    elem_col0 = node_col0
    ehdr = ["e", "y_e", "theta_e", "y_e+1", "theta_e+1", "fe1", "fe3"]
    for j, htxt in enumerate(ehdr):
        ws.cell(row=r, column=elem_col0 + j, value=htxt).font = FONT_SUBHEAD
        ws.cell(row=r, column=elem_col0 + j).fill = FILL_SUBHEAD
    r += 1
    elem_row0 = r
    EC = {name: elem_col0 + j for j, name in enumerate(["e", "ye", "the", "ye1", "the1", "fe1", "fe3"])}
    for e in range(n_elements):
        row = elem_row0 + e
        ws.cell(row=row, column=EC["e"], value=e)
        ye_c = a(final_y_row0 + e, y_col)
        the_c = a(final_theta_row0 + e, theta_col)
        ye1_c = a(final_y_row0 + e + 1, y_col)
        the1_c = a(final_theta_row0 + e + 1, theta_col)
        ws.cell(row=row, column=EC["ye"], value=f"={ye_c}")
        ws.cell(row=row, column=EC["the"], value=f"={the_c}")
        ws.cell(row=row, column=EC["ye1"], value=f"={ye1_c}")
        ws.cell(row=row, column=EC["the1"], value=f"={the1_c}")
        ye_r, the_r, ye1_r, the1_r = (a(row, EC["ye"]), a(row, EC["the"]),
                                       a(row, EC["ye1"]), a(row, EC["the1"]))
        fe1_f = f"={b_addr}*{ye_r}+{c_addr}*{the_r}-{b_addr}*{ye1_r}+{d_addr}*{the1_r}"
        fe3_f = f"={b_addr}*{ye_r}+{d_addr}*{the_r}-{b_addr}*{ye1_r}+{c_addr}*{the1_r}"
        ws.cell(row=row, column=EC["fe1"], value=fe1_f)
        ws.cell(row=row, column=EC["fe3"], value=fe3_f)
        for col in ("ye", "the", "ye1", "the1", "fe1", "fe3"):
            ws.cell(row=row, column=EC[col]).number_format = "0.000000"
    r = elem_row0 + n_elements + 1

    # M_node (nn rows), combining adjacent elements' end moments
    Mcol = elem_col0 + 7
    ws.cell(row=elem_row0 - 1, column=Mcol, value="M (kN*m)").font = FONT_SUBHEAD
    ws.cell(row=elem_row0 - 1, column=Mcol).fill = FILL_SUBHEAD
    M_row0 = elem_row0
    for i in range(nn):
        row = M_row0 + i
        if i == 0:
            f = f"=-{a(elem_row0, EC['fe1'])}"
        elif i == nn - 1:
            f = f"={a(elem_row0 + n_elements - 1, EC['fe3'])}"
        else:
            f = f"=({a(elem_row0 + i - 1, EC['fe3'])}-{a(elem_row0 + i, EC['fe1'])})/2"
        ws.cell(row=row, column=Mcol, value=f).number_format = "0.000"
    M_row_last = M_row0 + nn - 1
    r = max(r, M_row_last + 2)

    # --- outputs ------------------------------------------------------
    y_gl_cell = a(final_y_row0, y_col)
    theta_gl_cell = a(final_theta_row0, theta_col)
    ws.cell(row=r, column=col0, value="y_gl (mm)").font = Font(bold=True, size=10)
    ygl_addr = A(ws.title, r, col0 + 1)
    ws.cell(row=r, column=col0 + 1, value=f"={y_gl_cell}*1000").number_format = "0.000"
    ws.cell(row=r, column=col0 + 1).font = FONT_OUTPUT
    ws.cell(row=r, column=col0 + 1).fill = FILL_OUTPUT
    r += 1
    ws.cell(row=r, column=col0, value="theta_gl (deg)").font = Font(bold=True, size=10)
    thetagl_addr = A(ws.title, r, col0 + 1)
    ws.cell(row=r, column=col0 + 1, value=f"=DEGREES({theta_gl_cell})").number_format = "0.0000"
    ws.cell(row=r, column=col0 + 1).font = FONT_OUTPUT
    ws.cell(row=r, column=col0 + 1).fill = FILL_OUTPUT
    r += 1
    M_range = f"{a(M_row0, Mcol)}:{a(M_row_last, Mcol)}"
    ws.cell(row=r, column=col0, value="M_max (kN*m)").font = Font(bold=True, size=10)
    mmax_addr = A(ws.title, r, col0 + 1)
    _c = ws.cell(row=r, column=col0 + 1)
    _c.value = ArrayFormula(ref=plain(r, col0 + 1), text=f"=MAX(ABS({M_range}))")
    _c.number_format = "0.000"
    _c.font = FONT_OUTPUT
    _c.fill = FILL_OUTPUT
    r += 1

    return {
        "y_gl_mm": ygl_addr, "theta_gl_deg": thetagl_addr, "M_max": mmax_addr,
        "next_row": r + 1,
    }


# ===========================================================================
# STEP 7 -- CORROSION (built ahead of 5b/6, which both consume t_corroded)
# ===========================================================================
SH7 = "7-Corrosion"
ws = wb.create_sheet(SH7)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 7 -- CORROSION DESIGN (75-YR SACRIFICIAL THICKNESS)")
r = cite(ws, r, "corrosion.py -- galvanizing + sacrificial steel -> t_corroded feeding Step 6. Doc Section 7.")
r += 1
r = subhead(ws, r, "Inputs")
r = calc_row(ws, r, "Nominal wall thickness, t_nominal (Step 3)", f"={R['S3.tnom']}", "mm")
tnom7_row = r - 1
r = input_row(ws, r, "Zinc (HDG) thickness, t_zn", 100.0, "um"); tzn_row = r - 1
r = input_row(ws, r, "Zinc corrosion rate, r_zn", 4.0, "um/yr"); rzn_row = r - 1
r = input_row(ws, r, "Bare-steel corrosion rate, r_s", 20.0, "um/yr"); rs_row = r - 1
r = input_row(ws, r, "Design life", 75.0, "yr"); life_row = r - 1
lc = ws.cell(row=r, column=1, value="Both faces corrode (open/vented section)?")
lc.font = FONT_LABEL
bf_c = ws.cell(row=r, column=4, value="No")
bf_c.font = FONT_INPUT; bf_c.fill = FILL_INPUT; bf_c.border = BORDER_BOX
dv2 = DataValidation(type="list", formula1='"No,Yes"', allow_blank=False)
ws.add_data_validation(dv2); dv2.add(bf_c)
bothfaces_row = r
r += 1
r = input_row(ws, r, "Soil resistivity", 3000.0, "ohm-cm"); resis_row = r - 1
r = input_row(ws, r, "Soil pH", 7.0, "-"); ph_row = r - 1
r += 1

r = subhead(ws, r, "Calculated")
zlife_f = f"={a(tzn_row,4)}/{a(rzn_row,4)}"
r = calc_row(ws, r, "Zinc life", zlife_f, "yr"); zlife_row = r - 1
remain_f = f"=MAX({a(life_row,4)}-{a(zlife_row,4)},0)"
r = calc_row(ws, r, "Remaining design life after zinc consumed", remain_f, "yr"); remain_row = r - 1
tsac_f = f"={a(rs_row,4)}*{a(remain_row,4)}/1000"
r = calc_row(ws, r, "Sacrificial thickness, t_sac", tsac_f, "mm"); tsac_row = r - 1
tc_f = f"={a(tnom7_row,4)}-{a(tsac_row,4)}*IF({a(bothfaces_row,4)}=\"Yes\",2,1)"
r = calc_row(ws, r, "Corroded wall thickness, t_corroded", tc_f, "mm", bold_output=True); tc_row = r - 1
class_f = (f'=IF(OR({a(resis_row,4)}<1000,{a(ph_row,4)}<5.5),'
           f'"severely corrosive -- cathodic protection or epoxy required; confirm helical piles permitted per AC358",'
           f'IF(AND({a(resis_row,4)}>=2000,{a(ph_row,4)}>=5.5,{a(ph_row,4)}<=10,{a(life_row,4)}<=50),'
           f'"non-aggressive (screening basis)",'
           f'"moderately corrosive design case -- explicit sacrificial-thickness calc required"))')
r = calc_row(ws, r, "Corrosivity classification", class_f, "-", numfmt="General")

R.update({"C7.tc": A(SH7, tc_row, 4)})
ws.freeze_panes = "A4"
print("Step 7 (Corrosion) built")

# ===========================================================================
# STEP 5b -- p-y SOLVER SHEET
# ===========================================================================
SH5B = "5b-pySolver"
ws = wb.create_sheet(SH5B)
ws.sheet_view.showGridLines = False
set_col_widths(ws, {"A": 20, "B": 12, "C": 10})
for col in range(4, 60):
    ws.column_dimensions[get_column_letter(col)].width = 9

r = 1
r = sheet_header(ws, r, "STEP 5b -- p-y METHOD (FINAL LATERAL/MOMENT CHECK)", span=13)
r = cite(ws, r, "lateral_py.solve_py / pushover -- nonlinear beam-on-Matlock-soft-clay-Winkler-foundation. "
                 "Doc Section 5.3.", span=13)
r = note_row(ws, r,
             "No macros and no 'iterative calculation' setting needed. The Python solver's fixed-point "
             "relaxation loop is unrolled into sequential 'generation' blocks stacked down this sheet -- "
             "each generation's matrix solve depends only on the PREVIOUS generation's result, a one-way "
             "chain, not a circular reference. (An earlier version nested the matrix solve in a true "
             "self-referencing loop; tested in real Excel via COM automation, it was unreliable -- a "
             "transient error during the first iterative pass would latch permanently. This design avoids "
             "that failure mode entirely.) See the Cover sheet for the mesh-coarsening accuracy note.",
             span=13)
r += 1

svc_out = build_py_block(
    ws, r, 1, n_elements=20, n_generations=25,
    label="SERVICE CASE (n=20 elements, 25 unrolled iterations) -- eq. 5.3, serviceability check",
    V_ref=R["L1.Vsvc"], M_ref="-" + R["L1.Msvc"], D_ref=R["S3.Dshaft"], t_mm_ref=R["C7.tc"],
    L_ref=R["S3.tip"], su_ref=R["G2.Su"], gamma_ref=R["G2.gamma"], gammasub_ref=R["G2.gammasub"],
    gwt_ref=R["G2.gwt"], eps50_ref=R["G2.eps50"], frost_ref=R["G2.df"],
)
r = svc_out["next_row"] + 1

str_out = build_py_block(
    ws, r, 1, n_elements=20, n_generations=25,
    label="STRENGTH CASE (n=20 elements, 25 unrolled iterations) -- eq. 5.3, M_max feeds Step 6",
    V_ref=R["L1.V"], M_ref="-" + R["L1.M"], D_ref=R["S3.Dshaft"], t_mm_ref=R["C7.tc"],
    L_ref=R["S3.tip"], su_ref=R["G2.Su"], gamma_ref=R["G2.gamma"], gammasub_ref=R["G2.gammasub"],
    gwt_ref=R["G2.gwt"], eps50_ref=R["G2.eps50"], frost_ref=R["G2.df"],
)
r = str_out["next_row"] + 1

r = subhead(ws, r, "Serviceability acceptance -- eq. 5.3, doc Section 5.3", span=13)
r = input_row(ws, r, "Deflection limit, y_limit (A16 default -- confirm w/ owner)", 12.0, "mm")
ylim_row = r - 1
r = input_row(ws, r, "Rotation limit, theta_limit (A16 default -- confirm w/ owner)", 0.5, "deg")
thlim_row = r - 1
ypass_f = f"=ABS({svc_out['y_gl_mm']})<={a(ylim_row,4)}"
r = passfail_row(ws, r, "Deflection OK?", ypass_f); ypass_row = r - 1
thpass_f = f"=ABS({svc_out['theta_gl_deg']})<={a(thlim_row,4)}"
r = passfail_row(ws, r, "Rotation OK?", thpass_f); thpass_row = r - 1
svcpass_f = f"=AND({a(ypass_row,4)},{a(thpass_row,4)})"
r = passfail_row(ws, r, "Serviceability -- ALL PASS?", svcpass_f); svcpass_row = r - 1
yratio_f = f"=ABS({svc_out['y_gl_mm']})/{a(ylim_row,4)}"
r = calc_row(ws, r, "y_service / y_limit ratio (feeds Step 10 trigger)", yratio_f, "-", numfmt="0.00")
yratio_row = r - 1
r += 1

R.update({
    "P5.ysvc_mm": svc_out["y_gl_mm"], "P5.thsvc_deg": svc_out["theta_gl_deg"],
    "P5.ystr_mm": str_out["y_gl_mm"], "P5.thstr_deg": str_out["theta_gl_deg"],
    "P5.Mmax": str_out["M_max"], "P5.svcpass": A(SH5B, svcpass_row, 4),
    "P5.yratio": A(SH5B, yratio_row, 4),
})

r = subhead(ws, r, "Pushover sweep (n=10 elements per point) -- eq. 5.3-3, lambda_ult / lambda_required(=1.0)",
            span=13)
r = cite(ws, r, "Each lambda below scales the strength-case (V,M) and re-solves independently (own coarse "
                 "10-element mesh, 45 unrolled iterations -- more than Service/Strength's 25 because "
                 "convergence near collapse is much slower) -- a coarser mesh than Service/Strength since "
                 "only bracketing the ultimate-capacity factor matters here, not tight serviceability "
                 "precision. 'Collapse?' flags lambda where groundline deflection exceeds the 200 mm "
                 "runaway threshold used by lateral_py.pushover's y_collapse_m default.", span=13)
r += 1
push_summary_row = r
ws.cell(row=r, column=1, value="lambda").font = FONT_SUBHEAD
ws.cell(row=r, column=1).fill = FILL_SUBHEAD
ws.cell(row=r, column=2, value="y_gl (mm)").font = FONT_SUBHEAD
ws.cell(row=r, column=2).fill = FILL_SUBHEAD
ws.cell(row=r, column=3, value="Collapse?").font = FONT_SUBHEAD
ws.cell(row=r, column=3).fill = FILL_SUBHEAD
r += 1
push_summary_first = r
PUSHOVER_LAMBDAS = [1, 4, 8, 9, 10, 10.5]
push_block_start = r + len(PUSHOVER_LAMBDAS) + 2
push_outs = []
for lam in PUSHOVER_LAMBDAS:
    out = build_py_block(
        ws, push_block_start, 1, n_elements=10, n_generations=45,
        label=f"PUSHOVER lambda = {lam} (n=10 elements, 45 unrolled iterations)",
        V_ref=f"{lam}*{R['L1.V']}", M_ref=f"-{lam}*{R['L1.M']}", D_ref=R["S3.Dshaft"],
        t_mm_ref=R["C7.tc"], L_ref=R["S3.tip"], su_ref=R["G2.Su"], gamma_ref=R["G2.gamma"],
        gammasub_ref=R["G2.gammasub"], gwt_ref=R["G2.gwt"], eps50_ref=R["G2.eps50"], frost_ref=R["G2.df"],
    )
    push_outs.append(out)
    push_block_start = out["next_row"] + 1

for k, (lam, out) in enumerate(zip(PUSHOVER_LAMBDAS, push_outs)):
    row = push_summary_first + k
    ws.cell(row=row, column=1, value=lam).number_format = "0.0"
    ws.cell(row=row, column=2, value=f"={out['y_gl_mm']}").number_format = "0.0"
    ws.cell(row=row, column=3, value=f'=IF(ABS({out["y_gl_mm"]})>=200,"COLLAPSE","ok")')
push_summary_last = push_summary_first + len(PUSHOVER_LAMBDAS) - 1
lam_range = f"{a(push_summary_first,1)}:{a(push_summary_last,1)}"
collapse_range = f"{a(push_summary_first,3)}:{a(push_summary_last,3)}"
lambda_ult_f = ArrayFormula(
    ref=plain(push_summary_row, 6),
    text=f'=MAX(IF({collapse_range}="ok",{lam_range}))',
)
ws.cell(row=push_summary_row, column=5, value="lambda_ult =").font = Font(bold=True, size=10)
ws.cell(row=push_summary_row, column=6, value=lambda_ult_f)
ws.cell(row=push_summary_row, column=6).font = FONT_OUTPUT
ws.cell(row=push_summary_row, column=6).fill = FILL_OUTPUT
lamult_addr = a(push_summary_row, 6)
pushpass_f = f"={lamult_addr}>=2"
ws.cell(row=push_summary_row + 1, column=5, value="Pushover FS >= 2.0?").font = Font(bold=True, size=10)
pf_cell = ws.cell(row=push_summary_row + 1, column=6, value=pushpass_f)
pf_addr = a(push_summary_row + 1, 6)
ws.conditional_formatting.add(pf_addr, FormulaRule(formula=[f'{pf_addr}=TRUE'], fill=GREEN_FILL, font=GREEN_FONT))
ws.conditional_formatting.add(pf_addr, FormulaRule(formula=[f'{pf_addr}=FALSE'], fill=RED_FILL, font=RED_FONT))

R.update({"P5.lambda_ult": A(SH5B, push_summary_row, 6), "P5.pushpass": A(SH5B, push_summary_row + 1, 6)})
ws.freeze_panes = "A4"
print("Step 5b (p-y solver) built")

# ===========================================================================
# STEP 9 -- TORQUE (built ahead of Step 6, which references its T_install)
# ===========================================================================
SH9 = "9-Torque"
ws = wb.create_sheet(SH9)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 9 -- INSTALLATION TORQUE SPECIFICATION")
r = cite(ws, r, "torque.py -- dual termination criteria (torque AND depth); torque never waives depth. "
                 "Doc Section 9.")
r += 1
r = subhead(ws, r, "Inputs")
r = input_row(ws, r, "FS, axial (design basis)", 2.0, "-"); fsax_row = r - 1
r = calc_row(ws, r, "Compression demand, P_max (Step 4)", f"={R['A4.pmax']}", "kN")
pmax9_row = r - 1
r = calc_row(ws, r, "Frost-uplift demand, T_frost (Step 4)", f"={R['A4.tfrost']}", "kN")
tfrost9_row = r - 1
r = input_row(ws, r, "Torque coefficient, Kt [ASSUMED -- validate: AC358 report or site load test]",
              7.0, "1/m")
kt_row = r - 1
lc = ws.cell(row=r, column=1, value="Kt source")
lc.font = FONT_LABEL
ktsrc_c = ws.cell(row=r, column=4, value="ASSUMED -- validate before production")
ktsrc_c.font = FONT_INPUT; ktsrc_c.fill = FILL_INPUT; ktsrc_c.border = BORDER_BOX
dv3 = DataValidation(type="list", formula1='"AC358 report,site load test,ASSUMED -- validate before production"',
                      allow_blank=False)
ws.add_data_validation(dv3); dv3.add(ktsrc_c)
ktsrc_row = r
r += 1
r = input_row(ws, r, "Torque rating, T_rated", 15.0, "kN*m"); trated_row = r - 1
r = input_row(ws, r, "Anticipated/measured average torque, T_avg", 6.5, "kN*m"); tavg_row = r - 1
r = calc_row(ws, r, "Minimum tip depth required (Steps 3/5)", f"={R['S3.tip']}", "m")
tipmin_row = r - 1
r = input_row(ws, r, "Installed tip depth (as-built override)", None, "m")
tipactual_input_row = r - 1
tipactual_f = f"=IF({a(tipactual_input_row,4)}=\"\",{a(tipmin_row,4)},{a(tipactual_input_row,4)})"
r = calc_row(ws, r, "Installed tip depth (used)", tipactual_f, "m"); tipactual_row = r - 1
r += 1

r = subhead(ws, r, "Calculated -- doc eq. 9.2(1)")
preq_f = f"=MAX({a(fsax_row,4)}*{a(pmax9_row,4)},{a(fsax_row,4)}*{a(tfrost9_row,4)})"
r = calc_row(ws, r, "Required ultimate axial capacity, P_req", preq_f, "kN"); preq_row = r - 1
tmin_f = f"={a(preq_row,4)}/{a(kt_row,4)}"
r = calc_row(ws, r, "Minimum termination torque, T_min", tmin_f, "kN*m", bold_output=True); tmin_row = r - 1
r += 1

r = subhead(ws, r, "Dual termination criteria -- ALL THREE must be satisfied (doc Section 9.2)")
torqok_f = f"={a(tavg_row,4)}>={a(tmin_row,4)}"
r = passfail_row(ws, r, "(1) T_avg >= T_min?", torqok_f); torqok_row = r - 1
depthok_f = f"={a(tipactual_row,4)}>={a(tipmin_row,4)}"
r = passfail_row(ws, r, "(2) Installed depth >= minimum? (torque cannot waive this)", depthok_f)
depthok_row = r - 1
capok_f = f"={a(tavg_row,4)}<={a(trated_row,4)}"
r = passfail_row(ws, r, "(3) T_avg <= T_rated (never exceed rating)?", capok_f); capok_row = r - 1
allok_f = f"=AND({a(torqok_row,4)},{a(depthok_row,4)},{a(capok_row,4)})"
r = passfail_row(ws, r, "Step 9 termination criteria -- ALL PASS?", allok_f); allok9_row = r - 1

R.update({
    "T9.Tmin": A(SH9, tmin_row, 4), "T9.Tavg": A(SH9, tavg_row, 4), "T9.Trated": A(SH9, trated_row, 4),
    "T9.allpass": A(SH9, allok9_row, 4), "T9.ktassumed": A(SH9, ktsrc_row, 4),
})
ws.freeze_panes = "A4"
print("Step 9 (Torque) built")

# ===========================================================================
# STEP 6 -- STRUCTURAL (AISC 360, corroded section)
# ===========================================================================
SH6 = "6-Structural"
ws = wb.create_sheet(SH6)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 6 -- SHAFT STRUCTURAL CHECKS (AISC 360, CORRODED SECTION)")
r = cite(ws, r, "structural.py -- section compactness, H1 combined axial+flexure, shear, buckling trigger, "
                 "torsion cap. Doc Section 6. Uses t_corroded (Step 7), never nominal.")
r += 1
r = subhead(ws, r, "Section properties -- doc eq. 6.1")
r = calc_row(ws, r, "Shaft OD, D (Step 3)", f"={R['S3.Dshaft']}", "m"); d6_row = r - 1
r = calc_row(ws, r, "Corroded wall, t_c (Step 7)", f"={R['C7.tc']}/1000", "m"); tc6_row = r - 1
r = input_row(ws, r, "Steel yield strength, Fy", 345000.0, "kPa"); fy_row = r - 1
r = input_row(ws, r, "Elastic modulus, E", 200000000.0, "kPa"); e6_row = r - 1
di_f = f"={a(d6_row,4)}-2*{a(tc6_row,4)}"
r = calc_row(ws, r, "Inner diameter, d_i = D - 2*t_c", di_f, "m"); di_row = r - 1
s_f = f"=PI()/32*({a(d6_row,4)}^4-{a(di_row,4)}^4)/{a(d6_row,4)}"
r = calc_row(ws, r, "Elastic section modulus, S", s_f, "m^3", numfmt="0.0000000"); s_row = r - 1
z_f = f"=({a(d6_row,4)}^3-{a(di_row,4)}^3)/6"
r = calc_row(ws, r, "Plastic section modulus, Z", z_f, "m^3", numfmt="0.0000000"); z_row = r - 1
area_f = f"=PI()/4*({a(d6_row,4)}^2-{a(di_row,4)}^2)"
r = calc_row(ws, r, "Area, A", area_f, "m^2"); area_row = r - 1
lam_f = f"={a(d6_row,4)}/{a(tc6_row,4)}"
r = calc_row(ws, r, "Slenderness, D/t_c", lam_f, "-"); lam_row = r - 1
lamlim_f = f"=0.07*{a(e6_row,4)}/{a(fy_row,4)}"
r = calc_row(ws, r, "Compactness limit, 0.07*E/Fy (AISC Table B4.1b)", lamlim_f, "-"); lamlim_row = r - 1
compact_f = f"={a(lam_row,4)}<={a(lamlim_row,4)}"
r = passfail_row(ws, r, "Section compact?", compact_f); compact_row = r - 1
r += 1

r = subhead(ws, r, "H1 combined axial + flexure -- doc eq. 6.2")
r = input_row(ws, r, "phi_c (compression)", 0.90, "-"); phic_row = r - 1
r = input_row(ws, r, "phi_b (flexure)", 0.90, "-"); phib_row = r - 1
r = calc_row(ws, r, "Axial demand, Pr (Step 1 strength P)", f"={R['L1.P']}", "kN"); pr_row = r - 1
r = calc_row(ws, r, "Flexural demand, Mr (Step 5b strength M_max)", f"={R['P5.Mmax']}", "kN*m")
mr_row = r - 1
pc_f = f"={a(phic_row,4)}*{a(fy_row,4)}*{a(area_row,4)}"
r = calc_row(ws, r, "Pc = phi_c*Fy*A", pc_f, "kN"); pc_row = r - 1
mc_f = f"={a(phib_row,4)}*{a(fy_row,4)}*{a(z_row,4)}"
r = calc_row(ws, r, "Mc = phi_b*Fy*Z", mc_f, "kN*m"); mc_row = r - 1
prpc_f = f"={a(pr_row,4)}/{a(pc_row,4)}"
r = calc_row(ws, r, "Pr/Pc", prpc_f, "-", numfmt="0.0000"); prpc_row = r - 1
mrmc_f = f"={a(mr_row,4)}/{a(mc_row,4)}"
r = calc_row(ws, r, "Mr/Mc", mrmc_f, "-", numfmt="0.0000"); mrmc_row = r - 1
eq_f = f'=IF({a(prpc_row,4)}>=0.2,"H1-1a","H1-1b")'
r = calc_row(ws, r, "Governing equation", eq_f, "-", numfmt="General"); eq_row = r - 1
ratio_f = (f"=IF({a(prpc_row,4)}>=0.2,{a(prpc_row,4)}+(8/9)*{a(mrmc_row,4)},"
           f"{a(prpc_row,4)}/2+{a(mrmc_row,4)})")
r = calc_row(ws, r, "Interaction ratio", ratio_f, "-", numfmt="0.0000", bold_output=True); ratio_row = r - 1
h1pass_f = f"={a(ratio_row,4)}<=1"
r = passfail_row(ws, r, "H1 interaction passes (ratio<=1.0)?", h1pass_f); h1pass_row = r - 1
r += 1

r = subhead(ws, r, "Shear -- doc Section 6.2 (trivially satisfied for light poles, checked anyway)")
r = input_row(ws, r, "phi_v (shear)", 0.90, "-"); phiv_row = r - 1
r = calc_row(ws, r, "Shear demand, Vu (Step 1 strength V)", f"={R['L1.V']}", "kN"); vu_row = r - 1
vn_f = f"=0.6*{a(fy_row,4)}*({a(area_row,4)}/2)"
r = calc_row(ws, r, "Vn = 0.6*Fy*(A/2)", vn_f, "kN"); vn_row = r - 1
shearpass_f = f"={a(vu_row,4)}<={a(phiv_row,4)}*{a(vn_row,4)}"
r = passfail_row(ws, r, "Shear passes?", shearpass_f); shearpass_row = r - 1
r += 1

r = subhead(ws, r, "Column-buckling trigger -- doc Section 6.3 (advisory; run a separate Pn check if triggered)")
r = calc_row(ws, r, "Su (Step 2)", f"={R['G2.Su']}", "kPa"); su6_row = r - 1
r = input_row(ws, r, "Continuous weak layer thickness through which pile passes", 4.5, "m")
weaklayer_row = r - 1
lc = ws.cell(row=r, column=1, value="Shaft exposed or scoured?")
lc.font = FONT_LABEL
exp_c = ws.cell(row=r, column=4, value="No")
exp_c.font = FONT_INPUT; exp_c.fill = FILL_INPUT; exp_c.border = BORDER_BOX
dv4 = DataValidation(type="list", formula1='"No,Yes"', allow_blank=False)
ws.add_data_validation(dv4); dv4.add(exp_c)
exp_row = r
r += 1
trig_f = (f'=IF({a(exp_row,4)}="Yes",TRUE,'
          f'IF(AND({a(weaklayer_row,4)}>=1.5,{a(su6_row,4)}<25),TRUE,FALSE))')
r = passfail_row(ws, r, "Buckling check triggered? (TRUE = perform separate Pn check)", trig_f)
trig_row = r - 1
r += 1

r = subhead(ws, r, "Installation torsion cap -- doc eq. 6.4")
r = input_row(ws, r, "phi_T", 0.90, "-"); phit_row = r - 1
r = calc_row(ws, r, "T_install_max (Step 9 T_avg)", f"={R['T9.Tavg']}", "kN*m"); tinst_row = r - 1
r = calc_row(ws, r, "T_rated (Step 9)", f"={R['T9.Trated']}", "kN*m"); trated6_row = r - 1
torspass_f = f"={a(tinst_row,4)}<={a(phit_row,4)}*{a(trated6_row,4)}"
r = passfail_row(ws, r, "Torsion cap OK?", torspass_f); torspass_row = r - 1

R.update({
    "S6.compact": A(SH6, compact_row, 4), "S6.h1pass": A(SH6, h1pass_row, 4),
    "S6.shearpass": A(SH6, shearpass_row, 4), "S6.trigbuckling": A(SH6, trig_row, 4),
    "S6.torspass": A(SH6, torspass_row, 4), "S6.ratio": A(SH6, ratio_row, 4),
})
ws.freeze_panes = "A4"
print("Step 6 (Structural) built")

# ===========================================================================
# STEP 8 -- CONNECTION (bolts, weld, torsion slip)
# ===========================================================================
SH8 = "8-Connection"
ws = wb.create_sheet(SH8)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 8 -- POLE-TO-PILE CONNECTION DESIGN")
r = cite(ws, r, "connection.py -- anchor bolts (combined tension+standoff bending), fillet weld, "
                 "torsional slip. Doc Section 8.")
r += 1
r = subhead(ws, r, "Demands (Step 1 strength case)")
r = calc_row(ws, r, "M (overturning moment)", f"={R['L1.M']}", "kN*m"); m8_row = r - 1
r = calc_row(ws, r, "V (shear)", f"={R['L1.V']}", "kN"); v8_row = r - 1
r = calc_row(ws, r, "Tz (torsion, single-arm poles only)", f"={R['L1.Tz']}", "kN*m"); tz8_row = r - 1
r = calc_row(ws, r, "D_shaft (Step 3)", f"={R['S3.Dshaft']}", "m"); d8_row = r - 1
r = input_row(ws, r, "Net uplift on bolt group, P_uplift", 0.0, "kN"); puplift_row = r - 1
r += 1

r = subhead(ws, r, "8.1 Anchor bolt group geometry")
r = input_row(ws, r, "Number of bolts, n", 4, "-"); nbolts_row = r - 1
r = input_row(ws, r, "Bolt circle radius, r", 0.20, "m"); rbolt_row = r - 1
r = input_row(ws, r, "Bolt tensile stress area, Ab (per bolt)", 0.000314, "m^2"); ab_row = r - 1
r = input_row(ws, r, "Bolt Fu (e.g. F1554 Gr55)", 520000.0, "kPa"); fu_row = r - 1
r = input_row(ws, r, "Nominal bolt diameter (thread-root modulus not modeled -- see note)", 0.0254, "m")
boltd_row = r - 1
r = note_row(ws, r, "bolt_section_modulus_m3 uses the NOMINAL diameter as a stand-in for the (smaller) "
                     "thread-root diameter -- this overestimates the section modulus and therefore "
                     "UNDERESTIMATES bending stress. Acceptable for a preliminary screen only; replace with "
                     "the manufacturer's thread-root modulus before final design.")
r = input_row(ws, r, "Leveling-nut standoff", 0.0, "m"); standoff8_row = r - 1
r = input_row(ws, r, "phi_t (bolt tension, AISC 360 J3)", 0.75, "-"); phit8_row = r - 1
r += 1

r = subhead(ws, r, "Sum(ci^2) -- worst-case bending axis, n bolts evenly spaced on the circle")
bhdr = r
ws.cell(row=r, column=1, value="Bolt i").font = FONT_SUBHEAD
ws.cell(row=r, column=1).fill = FILL_SUBHEAD
ws.cell(row=r, column=4, value="ci^2 term (0 if i>=n)").font = FONT_SUBHEAD
ws.cell(row=r, column=4).fill = FILL_SUBHEAD
r += 1
bolt_term_rows = []
for i in range(12):
    ws.cell(row=r, column=1, value=i).font = FONT_LABEL
    term_f = (f"=IF({i}<{a(nbolts_row,4)},"
              f"({a(rbolt_row,4)}*COS(2*PI()*{i}/{a(nbolts_row,4)}))^2,0)")
    ws.cell(row=r, column=4, value=term_f).number_format = "0.000000"
    bolt_term_rows.append(r)
    r += 1
sumci2_f = f"=SUM({a(bolt_term_rows[0],4)}:{a(bolt_term_rows[-1],4)})"
r = calc_row(ws, r, "Sum(ci^2)", sumci2_f, "m^2", numfmt="0.000000"); sumci2_row = r - 1
r += 1

r = subhead(ws, r, "8.2 Combined bolt tension + standoff-bending check")
tbolt_f = f"={a(m8_row,4)}*{a(rbolt_row,4)}/{a(sumci2_row,4)}+{a(puplift_row,4)}/{a(nbolts_row,4)}"
r = calc_row(ws, r, "Bolt tension demand, T_bolt = M*c_max/Sum(ci^2) + P_uplift/n", tbolt_f, "kN")
tbolt_row = r - 1
sigt_f = f"={a(tbolt_row,4)}/{a(ab_row,4)}"
r = calc_row(ws, r, "Tension stress, sigma_tension = T_bolt/Ab", sigt_f, "kPa"); sigt_row = r - 1
bendmom_f = (f"=IF({a(standoff8_row,4)}<={a(boltd_row,4)},0,"
             f"({a(v8_row,4)}/{a(nbolts_row,4)})*{a(standoff8_row,4)}/2)")
r = calc_row(ws, r, "Standoff bending moment, M_b (0 if standoff <= bolt diam)", bendmom_f, "kN*m")
bendmom_row = r - 1
zbolt_f = f"=PI()*{a(boltd_row,4)}^3/32"
r = calc_row(ws, r, "Bolt section modulus, Z_bolt (solid shank)", zbolt_f, "m^3", numfmt="0.00000000")
zbolt_row = r - 1
sigb_f = f"=IF({a(zbolt_row,4)}>0,{a(bendmom_row,4)}/{a(zbolt_row,4)},0)"
r = calc_row(ws, r, "Bending stress, sigma_bending", sigb_f, "kPa"); sigb_row = r - 1
sigtot_f = f"={a(sigt_row,4)}+{a(sigb_row,4)}"
r = calc_row(ws, r, "Combined stress, sigma_total", sigtot_f, "kPa", bold_output=True); sigtot_row = r - 1
sigcap_f = f"={a(phit8_row,4)}*0.75*{a(fu_row,4)}"
r = calc_row(ws, r, "Capacity, sigma_capacity = phi_t*0.75*Fu", sigcap_f, "kPa"); sigcap_row = r - 1
boltpass_f = f"={a(sigtot_row,4)}<={a(sigcap_row,4)}"
r = passfail_row(ws, r, "Bolt combined-stress check passes?", boltpass_f); boltpass_row = r - 1
r += 1

r = subhead(ws, r, "8.3 Fillet weld (shaft-to-cap-plate, all-around, treated as a line) -- doc eq. 8.3")
r = input_row(ws, r, "Fillet weld size, a_w", 0.008, "m"); fillet_row = r - 1
r = input_row(ws, r, "Electrode strength, FEXX", 490000.0, "kPa"); fexx_row = r - 1
r = input_row(ws, r, "phi_w (weld, AISC 360 J2)", 0.75, "-"); phiw_row = r - 1
rw_f = f"={a(d8_row,4)}/2"
r = calc_row(ws, r, "Weld-line radius, r = D/2", rw_f, "m"); rw_row = r - 1
sw_f = f"=PI()*{a(rw_row,4)}^2"
r = calc_row(ws, r, "Weld section modulus, Sw = pi*r^2", sw_f, "m^2", numfmt="0.000000"); sw_row = r - 1
wm_f = f"={a(m8_row,4)}/{a(sw_row,4)}"
r = calc_row(ws, r, "w_M = M/Sw", wm_f, "kN/m"); wm_row = r - 1
wv_f = f"={a(v8_row,4)}/(PI()*{a(d8_row,4)})"
r = calc_row(ws, r, "w_V = V/(pi*D)", wv_f, "kN/m"); wv_row = r - 1
wr_f = f"=SQRT({a(wm_row,4)}^2+{a(wv_row,4)}^2)"
r = calc_row(ws, r, "Resultant weld demand, w_r", wr_f, "kN/m", bold_output=True); wr_row = r - 1
wcap_f = f"={a(phiw_row,4)}*0.6*{a(fexx_row,4)}*0.707*{a(fillet_row,4)}"
r = calc_row(ws, r, "Weld capacity", wcap_f, "kN/m"); wcap_row = r - 1
weldpass_f = f"={a(wr_row,4)}<={a(wcap_row,4)}"
r = passfail_row(ws, r, "Weld check passes?", weldpass_f); weldpass_row = r - 1
r += 1

r = subhead(ws, r, "8.4 Torsional slip resistance (single-arm poles only) -- doc Section 8.4")
r = input_row(ws, r, "Pile torsional slip capacity [ASSUMED -- validate]", 15.0, "kN*m")
piletors_row = r - 1
r = input_row(ws, r, "Required factor on Tz", 3.0, "-"); torsfactor_row = r - 1
torspass8_f = f"={a(piletors_row,4)}>={a(torsfactor_row,4)}*{a(tz8_row,4)}"
r = passfail_row(ws, r, "Torsional slip resistance OK (capacity >= factor*Tz)?", torspass8_f)
torspass8_row = r - 1
r = note_row(ws, r, "If Tz = 0 (no mast-arm eccentricity, e.g. twin/symmetric fixtures), this check is "
                     "automatically satisfied -- confirm arm eccentricity on Step 1 reflects the actual "
                     "luminaire configuration.")

R.update({
    "C8.boltpass": A(SH8, boltpass_row, 4), "C8.weldpass": A(SH8, weldpass_row, 4),
    "C8.torspass": A(SH8, torspass8_row, 4),
})
ws.freeze_panes = "A4"
print("Step 8 (Connection) built")

# ===========================================================================
# STEP 10 -- QA/QC AND LOAD-TEST TRIGGERS
# ===========================================================================
SH10 = "10-QAQC"
ws = wb.create_sheet(SH10)
set_col_widths(ws, DEFAULT_WIDTHS)
r = 1
r = sheet_header(ws, r, "STEP 10 -- QA/QC AND LOAD-TEST TRIGGERS")
r = cite(ws, r, "qaqc.py -- rule engine reading flags from every prior step. Doc Section 10.2. "
                 "Not a calculation -- a trigger table for ASTM D1143/D3689/D3966.")
r += 1
r = subhead(ws, r, "Trigger inputs (mostly linked from prior steps)")
kt_ac358_f = f'={R["T9.ktassumed"]}="AC358 report"'
r = passfail_row(ws, r, "Kt established by AC358 report?", kt_ac358_f); ktac358_row = r - 1
r = input_row(ws, r, "Frost-uplift design relies on assumed tau_ad/sleeve performance?", "Yes")
frostassumed_input_row = r - 1
frostassumed_f = f'={a(frostassumed_input_row,4)}="Yes"'
r = calc_row(ws, r, "  (as boolean)", frostassumed_f, "-", numfmt="General"); frostassumed_row = r - 1
pycpt_f = f"={R['G2.labcpt']}"
r = calc_row(ws, r, "p-y params from lab/CPT (Step 2)?", pycpt_f, "-", numfmt="General"); pycpt_row = r - 1
r = calc_row(ws, r, "Service deflection / limit ratio (Step 5b)", f"={R['P5.yratio']}", "-", numfmt="0.00")
yratio10_row = r - 1
r = input_row(ws, r, "Number of production piles", 40, "-"); npiles_row = r - 1
lc = ws.cell(row=r, column=1, value="Variable or unfamiliar soils across the site?")
lc.font = FONT_LABEL
varsoil_c = ws.cell(row=r, column=4, value="No")
varsoil_c.font = FONT_INPUT; varsoil_c.fill = FILL_INPUT; varsoil_c.border = BORDER_BOX
dv5 = DataValidation(type="list", formula1='"No,Yes"', allow_blank=False)
ws.add_data_validation(dv5); dv5.add(varsoil_c)
varsoil_row = r
r += 1
lc = ws.cell(row=r, column=1, value="Owner/code official requires a load test?")
lc.font = FONT_LABEL
ownerreq_c = ws.cell(row=r, column=4, value="No")
ownerreq_c.font = FONT_INPUT; ownerreq_c.fill = FILL_INPUT; ownerreq_c.border = BORDER_BOX
dv6 = DataValidation(type="list", formula1='"No,Yes"', allow_blank=False)
ws.add_data_validation(dv6); dv6.add(ownerreq_c)
ownerreq_row = r
r += 1

r = subhead(ws, r, "Trigger evaluation -- doc Section 10.2")
t1_f = f"=NOT({a(ktac358_row,4)})"
r = passfail_row(ws, r, "D1143/D3689 triggered? (Kt not AC358-verified)", t1_f); t1_row = r - 1
t2_f = f"={a(frostassumed_row,4)}"
r = passfail_row(ws, r, "D3689 triggered? (frost-uplift relies on assumed values)", t2_f); t2_row = r - 1
t3_f = f'=OR({a(yratio10_row,4)}>0.75,NOT({a(pycpt_row,4)}))'
r = passfail_row(ws, r, "D3966 triggered? (low lateral margin OR p-y not lab/CPT)", t3_f); t3_row = r - 1
t4_f = f'=OR({a(npiles_row,4)}>50,{a(varsoil_row,4)}="Yes")'
r = passfail_row(ws, r, "D3689+D3966 triggered? (>50 piles OR variable soils)", t4_f); t4_row = r - 1
t5_f = f'={a(ownerreq_row,4)}="Yes"'
r = passfail_row(ws, r, "Owner/code-official-directed test triggered?", t5_f); t5_row = r - 1
anytrig_f = f"=OR({a(t1_row,4)},{a(t2_row,4)},{a(t3_row,4)},{a(t4_row,4)},{a(t5_row,4)})"
r = passfail_row(ws, r, "ANY load test triggered?", anytrig_f); anytrig_row = r - 1
r += 1

r = subhead(ws, r, "Baseline QA items (always applicable)")
for item in [
    "Mill certs for shaft/helix steel and HDG on file",
    "AC358 evaluation report (or equivalent product data) on file",
    "Calibrated torque indicator certificate <= 12 months old",
    "Installation logs (torque-vs-depth, per pile) reviewed and signed by the engineer",
    "Anchor-bolt template verification completed before pole delivery",
]:
    ws.cell(row=r, column=1, value="[ ] " + item).font = Font(size=10)
    r += 1

R.update({"Q10.anytrig": A(SH10, anytrig_row, 4)})
ws.freeze_panes = "A4"
print("Step 10 (QAQC) built")

# ===========================================================================
# SUMMARY / PASS-FAIL DASHBOARD
# ===========================================================================
SHS = "Summary"
ws = wb.create_sheet(SHS)
set_col_widths(ws, {"A": 10, "B": 55, "C": 16, "D": 40})
r = 1
r = sheet_header(ws, r, "PASS/FAIL ROLLUP -- ALL STEPS", span=4)
r = cite(ws, r, "Pulls the governing pass/fail cell from each step sheet. GREEN = pass, RED = fail. "
                 "This does not replace engineer review of every assumption cell (pale orange) on each sheet.",
         span=4)
r += 1
hdr_row = r
for j, htxt in enumerate(["Step", "Check", "Result", "Notes"], start=1):
    ws.cell(row=hdr_row, column=j, value=htxt).font = FONT_SUBHEAD
    ws.cell(row=hdr_row, column=j).fill = FILL_SUBHEAD
r += 1

rows = [
    ("3", "Embedment screens (frost clearance + deep-foundation mode)", R["S3.allpass"], ""),
    ("4", "Axial compression, FS>=2.0", R["A4.passc"], ""),
    ("4", "Axial tension vs. frost uplift, FS>=2.0", R["A4.passt"], ""),
    ("5a", "Broms preliminary lateral, FS>=2.0", R["B5.pass"], "Cross-check only, doc 5.2"),
    ("5b", "p-y serviceability (deflection + rotation)", R["P5.svcpass"], ""),
    ("5b", "p-y pushover, FS>=2.0", R["P5.pushpass"], f"lambda_ult = see 5b-pySolver!{R['P5.lambda_ult']}"),
    ("6", "Section compactness", R["S6.compact"], ""),
    ("6", "H1 combined axial+flexure", R["S6.h1pass"], ""),
    ("6", "Shear", R["S6.shearpass"], ""),
    ("6", "Installation torsion cap", R["S6.torspass"], ""),
    ("8", "Bolt combined tension+bending", R["C8.boltpass"], "Uses nominal-diameter Z_bolt -- see Step 8 note"),
    ("8", "Fillet weld", R["C8.weldpass"], ""),
    ("8", "Torsional slip (single-arm poles)", R["C8.torspass"], ""),
    ("9", "Installation torque -- dual termination criteria", R["T9.allpass"], ""),
]
for step, label, addr, note in rows:
    ws.cell(row=r, column=1, value=step).font = Font(bold=True, size=10)
    ws.cell(row=r, column=2, value=label).font = FONT_LABEL
    vc = ws.cell(row=r, column=3, value=f"={addr}")
    vc.font = FONT_OUTPUT
    vcaddr = a(r, 3)
    ws.conditional_formatting.add(vcaddr, FormulaRule(formula=[f'{vcaddr}=TRUE'], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(vcaddr, FormulaRule(formula=[f'{vcaddr}=FALSE'], fill=RED_FILL, font=RED_FONT))
    ws.cell(row=r, column=4, value=note).font = FONT_CITE
    r += 1
r += 1
allpass_range = f"C{hdr_row+1}:C{r-2}"
r = subhead(ws, r, "Overall", span=4)
overall_f = f"=AND({allpass_range})"
ws.cell(row=r, column=2, value="ALL CHECKS PASS?").font = Font(bold=True, size=11)
oc = ws.cell(row=r, column=3, value=overall_f)
oc.font = Font(bold=True, size=12)
oaddr = a(r, 3)
ws.conditional_formatting.add(oaddr, FormulaRule(formula=[f'{oaddr}=TRUE'], fill=GREEN_FILL, font=GREEN_FONT))
ws.conditional_formatting.add(oaddr, FormulaRule(formula=[f'{oaddr}=FALSE'], fill=RED_FILL, font=RED_FONT))
r += 2
r = subhead(ws, r, "Step 10 -- QA/QC", span=4)
ws.cell(row=r, column=2, value="Any full-scale load test triggered?").font = Font(bold=True, size=10)
qc = ws.cell(row=r, column=3, value=f"={R['Q10.anytrig']}")
qc.font = FONT_OUTPUT
qaddr = a(r, 3)
ws.conditional_formatting.add(qaddr, FormulaRule(formula=[f'{qaddr}=TRUE'], fill=RED_FILL, font=RED_FONT))
ws.conditional_formatting.add(qaddr, FormulaRule(formula=[f'{qaddr}=FALSE'], fill=GREEN_FILL, font=GREEN_FONT))
ws.cell(row=r, column=4, value="See 10-QAQC for which test(s) and why.").font = FONT_CITE

ws.freeze_panes = "A4"
print("Summary sheet built")

# ===========================================================================
# Tab order + workbook calc settings
# ===========================================================================
order = ["Cover", SH1, SH2, SH3, SH4, SH5A, SH5B, SH6, SH7, SH8, SH9, SH10, SHS]
wb._sheets = [wb[name] for name in order]
for name in order:
    wb[name].sheet_view.tabSelected = False
wb.active = 0

wb.calculation.iterate = False
wb.calculation.fullCalcOnLoad = True

OUT_PATH = "HelicalPileDesign_CalcSheet.xlsx"
wb.save(OUT_PATH)
print(f"saved {OUT_PATH}")













